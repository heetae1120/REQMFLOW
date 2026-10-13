import io
import tempfile
import unittest
import zipfile
from unittest.mock import patch
from pathlib import Path
from openpyxl import load_workbook
from reqm_local.service import Operations
from reqm_local.files import workbook_bytes,WEKEEP_REQUEST_COLUMNS
from reqm_local.tracking_store import encode
from reqm_local.workspace_cloud import export_workspace,import_workspace,workspace_digest
from tests.test_tracking import SOURCE_HEADERS,ORDER,REMOTE,DAY,REFERENCE


class DailyB2CTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.folder=Path(self.temp.name)
        self.s=Operations(self.folder/'data',REFERENCE)
        self.s.settings['profiles']=[{'channel':'오늘의집','required':['주문번호','상품주문번호','상품명','수량']}]

    def tearDown(self):
        self.s.close();self.temp.cleanup()

    def request(self,number='A',channel='오늘의집'):
        path=self.folder/(number+'.xlsx')
        path.write_bytes(workbook_bytes(SOURCE_HEADERS,[[number,'LINE-'+number,ORDER['product'],ORDER['option'],1,1000,ORDER['recipient'],ORDER['phone'],ORDER['postcode'],ORDER['address'],0]]))
        self.s.import_files([path],channel_override=channel)
        order=next(o for o in self.s.orders() if o['data']['order_no']==number)
        self.s.set_mapping(order['id'],[dict(code='QWC-Q3100S-BK',logistics_code='QWC-Q3100S-BK',name='변환품목',quantity=1,unit_amount=0,warehouse='300',customer='DEMO')])
        batch=self.s.request([order['id']],DAY)
        return batch,order['id']

    def remote(self,number,channel='오늘의집',tracking=REMOTE['tracking']):
        return {**REMOTE,'order_no':number,'line_no':'','channel':channel,'tracking':tracking}

    def test_request_stores_without_download_and_registers_once(self):
        job,oid=self.request()
        self.assertEqual(len(self.s.daily_b2c_jobs(DAY)),1)
        self.s.daily_b2c_jobs(DAY)
        self.assertEqual(len(self.s.tracking_rows(job)),1)
        self.assertEqual(self.s.tracking_rows(job)[0]['order_id'],oid)
        self.assertIsNotNone(self.s.db.execute('SELECT content FROM artifacts WHERE id=?',(job,)).fetchone())
        self.assertEqual(self.s.orders()[0]['state'],'출고 요청')

    def test_wekeep_export_resolves_old_cached_names_without_changing_orders(self):
        job,oid=self.request()
        order=next(o for o in self.s.orders() if o['id']==oid)
        components=order['components']
        components[0]['name']='[리큐엠] 예전 상품명'
        with self.s.db:
            self.s.db.execute('UPDATE orders SET components=?,state=? WHERE id=?',
                (encode(components),'출고 준비',oid))
        expected=self.s.wekeep_output_name(components[0])
        self.assertNotIn('리큐엠',expected)
        self.assertNotEqual(expected,'예전 상품명')
        # Old tracking jobs must also resolve the current item name on download.
        target=self.folder/'old-job.xlsx'
        self.s.export_wekeep_jobs([job],target,True)
        book=load_workbook(target)
        self.assertEqual(book.active['C2'].value,expected);book.close()
        with self.s.db:self.s.db.execute('DELETE FROM request_lines WHERE request_id=?',(job,))
        new_job=self.s.request([oid],DAY)
        content=self.s.db.execute('SELECT content FROM artifacts WHERE id=?',(new_job,)).fetchone()[0]
        book=load_workbook(io.BytesIO(content))
        self.assertEqual(book.active['C2'].value,expected);book.close()
        self.assertEqual(self.s.orders()[0]['components'][0]['name'],'[리큐엠] 예전 상품명')

    def test_brand_prefix_removal_preserves_product_details(self):
        from reqm_local.output_names import without_brand_prefix
        self.assertEqual(without_brand_prefix('[리큐엠]C.class K-CC150-실리콘케이블 블루(100W)'),
                         'C.class K-CC150-실리콘케이블 블루(100W)')
        self.assertEqual(without_brand_prefix('리큐엠-UV살균기 무선충전 시계_블랙'),
                         'UV살균기 무선충전 시계_블랙')
        self.assertEqual(without_brand_prefix('[1.5M] 실리콘케이블'),'[1.5M] 실리콘케이블')

    def test_batch_auto_confirms_and_keeps_shipping_state_and_exports(self):
        a,oid=self.request('A');b,_=self.request('B','무신사')
        expected={j:encode(self.s.active_tracking_rows(j)) for j in (a,b)}
        self.s.apply_daily_tracking([a,b],[self.remote('A'),self.remote('B','무신사')],DAY,DAY,expected)
        self.assertTrue(all(self.s.tracking_rows(j)[0]['state']=='confirmed' for j in (a,b)))
        self.assertTrue(all(o['state']=='출고 요청' for o in self.s.orders()))
        self.assertEqual(self.s.db.execute('SELECT COUNT(*) FROM shipments').fetchone()[0],0)
        self.assertEqual(next(o for o in self.s.orders() if o['id']==oid)['data']['tracking'],REMOTE['tracking'])
        target=self.folder/'wekeep.xlsx';self.s.export_wekeep_jobs([a,b],target,True,True)
        book=load_workbook(target);self.assertEqual([c.value for c in book.active[1]],WEKEEP_REQUEST_COLUMNS)
        self.assertEqual(book.active['J2'].value,REMOTE['tracking']);self.assertEqual(book.active['J2'].data_type,'s');book.close()
        archive=self.folder/'sources.zip';self.s.export_daily_sources([a,b],archive)
        with zipfile.ZipFile(archive) as z:
            self.assertEqual(len(z.namelist()),2)
            book=load_workbook(io.BytesIO(z.read(z.namelist()[0])))
            self.assertEqual([c.value for c in book.active[1]][:len(SOURCE_HEADERS)],SOURCE_HEADERS)
            self.assertEqual(book.active.cell(2,len(SOURCE_HEADERS)+1).value,REMOTE['tracking']);book.close()
        state=export_workspace(self.s);import_workspace(self.s,state)
        self.assertEqual(workspace_digest(export_workspace(self.s)),workspace_digest(state))

    def test_cancelled_line_and_concurrent_changes_are_not_applied(self):
        job,oid=self.request()
        expected={job:encode(self.s.active_tracking_rows(job))}
        self.s.delete_orders([oid])
        self.assertEqual(self.s.daily_b2c_jobs(DAY),[])
        with self.assertRaisesRegex(ValueError,'변경'):
            self.s.apply_daily_tracking([job],[self.remote('A')],DAY,DAY,expected)
        self.assertEqual(self.s.db.execute('SELECT COUNT(*) FROM tracking_queries').fetchone()[0],0)

    def test_duplicate_order_numbers_across_markets_are_reviewed(self):
        a,_=self.request('A');b,_=self.request('B','무신사')
        data=self.s.tracking_rows(b)[0]
        changed={**data['data'],'order_no':'A'}
        with self.s.db:self.s.db.execute('UPDATE tracking_rows SET data=? WHERE id=?',(encode(changed),data['id']))
        self.s.apply_daily_tracking([a,b],[self.remote('A',channel='')],DAY,DAY)
        self.assertEqual([self.s.tracking_rows(j)[0]['state'] for j in (a,b)],['review','review'])

    def test_conflicting_order_invoice_rolls_back_entire_batch(self):
        a,_=self.request('A');b,oid=self.request('B')
        order=next(o for o in self.s.orders() if o['id']==oid)
        with self.s.db:self.s.db.execute('UPDATE orders SET data=? WHERE id=?',(encode({**order['data'],'tracking':'999999999999'}),oid))
        with self.assertRaises(ValueError):self.s.apply_daily_tracking([a,b],[self.remote('A'),self.remote('B')],DAY,DAY)
        self.assertEqual([self.s.tracking_rows(j)[0]['state'] for j in (a,b)],['pending','pending'])
        self.assertEqual(self.s.db.execute('SELECT COUNT(*) FROM tracking_queries').fetchone()[0],0)

    def test_requery_preserves_confirmed_invoice(self):
        job,_=self.request()
        self.s.apply_daily_tracking([job],[self.remote('A')],DAY,DAY)
        self.s.apply_daily_tracking([job],[self.remote('A',tracking='999999999999')],DAY,DAY)
        row=self.s.tracking_rows(job)[0]
        self.assertEqual(row['tracking'],REMOTE['tracking']);self.assertIn('확정 송장 유지',row['reason'])

    def test_screen_fetches_b2c_once_for_all_selected_jobs(self):
        import os
        import tkinter as tk
        from reqm_local.desktop import Desktop
        a,_=self.request('A');b,_=self.request('B','무신사')
        root=tk.Tk();root.withdraw()
        try:
            with patch.dict(os.environ,{'REQM_SELF_TEST':'1'}):app=Desktop(root,self.s)
            panel=app.tracking_panel;panel.day.set(DAY);panel.refresh();root.update_idletasks()
            self.assertEqual(set(panel.selected_jobs()),{a,b})
            with patch('reqm_local.tracking_ui.fetch_tracking',return_value=[self.remote('A'),self.remote('B','무신사')]) as fetch,patch('reqm_local.tracking_ui.threading.Thread') as thread:
                thread.return_value.start.side_effect=lambda:thread.call_args.kwargs['target']()
                panel.fetch();panel.poll()
            fetch.assert_called_once()
            self.assertEqual(fetch.call_args.args[3],'b2c')
            self.assertTrue(all(self.s.tracking_rows(j)[0]['state']=='confirmed' for j in (a,b)))
            self.assertEqual(len(panel.row_tree.get_children()),2)
            self.assertEqual(panel.row_tree['columns'][1],'송장번호')
        finally:root.destroy()
