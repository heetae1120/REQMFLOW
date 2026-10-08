import io
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from reqm_local.files import workbook_bytes
from reqm_local.service import Operations
from reqm_local.tracking import HEADERS, reconcile, manual_invoice
from reqm_local.tracking_store import encode
from reqm_local.wekeep_tracking import parse_remote_rows, read_tracking_file, detail_rows
from reqm_local.workspace_cloud import export_workspace, import_workspace, workspace_digest

REFERENCE = Path(__file__).resolve().parents[1] / 'supabase/ecount_migration/data'
DAY = '2026-10-08'
ORDER = dict(order_no='0001-AB', line_no='L1', channel='오늘의집', account='기본',
             recipient='테스트 수령인', phone='01012345678', postcode='00123', address='서울 테스트로 1 101호',
             product='테스트상품', option='기본', quantity='1', memo='문 앞')
REMOTE = {**ORDER, 'tracking': '001234567890', 'additional_tracking': ''}
SOURCE_HEADERS = ['주문번호', '상품주문번호', '상품명', '옵션정보', '수량', '최종 상품별 총 주문금액',
                  '수취인명', '수취인연락처1', '우편번호', '통합배송지', '배송비 합계']


class MatcherTests(unittest.TestCase):
    def test_matching_normalizes_spacing_and_phone_but_preserves_order_identity(self):
        result = reconcile([ORDER], [{**REMOTE, 'recipient': '테스트수령인', 'phone': '+82-10-1234-5678'}])[0]
        self.assertEqual((result['state'], result['tracking']), ('matched', REMOTE['tracking']))
        result = reconcile([ORDER], [{**REMOTE, 'order_no': '0001AB'}])[0]
        self.assertEqual(result['state'], 'not_found')

    def test_wrong_or_missing_delivery_data_never_fills_invoice(self):
        for key in ('recipient', 'phone', 'postcode', 'address'):
            for value in ('', '다른정보'):
                result = reconcile([ORDER], [{**REMOTE, key: value}])[0]
                self.assertEqual((result['state'], result['tracking']), ('review', ''))

    def test_ambiguous_identity_additional_invoice_and_pending(self):
        self.assertEqual(reconcile([ORDER], [REMOTE, {**REMOTE, 'recipient': '다른사람'}])[0]['state'], 'review')
        self.assertEqual(reconcile([ORDER], [{**REMOTE, 'additional_tracking': '999999999999'}])[0]['state'], 'review')
        self.assertEqual(reconcile([ORDER], [{**REMOTE, 'tracking': '-'}])[0]['state'], 'pending')

    def test_delivery_peers_are_not_automatically_treated_as_combined_packages(self):
        second = {**ORDER, 'order_no': 'O2'}
        results = reconcile([ORDER, second], [REMOTE])
        self.assertEqual([r['state'] for r in results], ['matched', 'not_found'])

    def test_manual_number_validation(self):
        self.assertEqual(manual_invoice('0012-3456-7890'), '001234567890')
        for value in ('123', '123456789a', '12345678;99999999', '=1234567890'):
            with self.assertRaises(ValueError): manual_invoice(value)

    def test_remote_file_and_detail_parsers(self):
        headers = ['판매처주문번호', '수령자', '핸드폰', '우편번호', '주소', '상세주소', '송장번호']
        rows = parse_remote_rows(headers, [['O1', '홍길동', '01012345678', '00123', '서울', '101호', '1234567890']])
        self.assertEqual(rows[0]['address'], '서울 101호')

    def test_detail_identity_and_provider_confirmed_combined_package(self):
        api = {'result': True, 'data': {
            'orderDetail': {'orderNo': 'INTERNAL', 'saleChannelNo': 'B2C', 'channelOrderNo': ORDER['order_no'],
                            'recipient': ORDER['recipient'], 'recipientMobile': ORDER['phone'],
                            'deliveryZipcode': ORDER['postcode'], 'deliveryAddress': '서울 테스트로 1',
                            'deliveryAddressDetail': '101호', 'groupOrderYn': 'Y'},
            'deliveryDetail': {'trackingNo': REMOTE['tracking']},
            'orderItemList': [{'channelOrderNo': ORDER['order_no']}, {'channelOrderNo': 'O2'}]}}
        rows = detail_rows(REMOTE, api, 'INTERNAL', 'B2C')
        self.assertEqual([r['order_no'] for r in rows], [ORDER['order_no'], 'O2'])
        self.assertEqual(rows[0]['address'], ORDER['address'])
        self.assertTrue(rows[1]['provider_grouped'])
        self.assertEqual([r['state'] for r in reconcile([ORDER, {**ORDER, 'order_no': 'O2'}], rows)], ['matched', 'matched'])
        with self.assertRaises(RuntimeError): detail_rows(REMOTE, api, 'WRONG', 'B2C')
        with self.assertRaises(RuntimeError): detail_rows(REMOTE, api, 'INTERNAL', 'OTHER')


class TrackingStoreTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.folder = Path(self.temp.name)
        self.s = Operations(self.folder / 'data', REFERENCE)
        self.path = self.folder / 'orders.xlsx'
        self.path.write_bytes(workbook_bytes(SOURCE_HEADERS, [[ORDER['order_no'], ORDER['line_no'], ORDER['product'],
            ORDER['option'], 1, 1000, ORDER['recipient'], ORDER['phone'], ORDER['postcode'], ORDER['address'], 0]]))

    def tearDown(self):
        self.s.close(); self.temp.cleanup()

    def job(self):
        return self.s.import_tracking_files([self.path], DAY, channel_override='오늘의집')[0][0]

    def query(self, job, remote=None):
        self.s.apply_tracking_query(job, [REMOTE] if remote is None else remote, DAY, DAY)

    def test_durable_jobs_deduplicate_and_keep_original(self):
        job = self.job()
        jobs, dup = self.s.import_tracking_files([self.path], DAY, channel_override='오늘의집')
        self.assertEqual((jobs, dup), ([job], 1))
        self.query(job)
        self.s.close(); self.s = Operations(self.folder / 'data', REFERENCE)
        self.assertEqual(self.s.tracking_rows(job)[0]['tracking'], REMOTE['tracking'])
        self.assertEqual(self.s.db.execute('SELECT source_content FROM tracking_jobs').fetchone()[0], self.path.read_bytes())

    def test_confirm_links_order_without_creating_actual_shipment(self):
        self.s.import_files([self.path], channel_override='오늘의집')
        job = self.job(); row = self.s.tracking_rows(job)[0]
        self.assertIsNotNone(row['order_id'])
        state = self.s.orders()[0]['state']
        self.query(job)
        self.s.confirm_tracking(job, [row['id']])
        self.assertEqual(self.s.orders()[0]['data']['tracking'], REMOTE['tracking'])
        self.assertEqual(self.s.orders()[0]['state'], state)
        self.assertEqual(self.s.db.execute('SELECT COUNT(*) FROM shipments').fetchone()[0], 0)
        self.s._sync_source_tracking([row['order_id']])
        self.assertEqual(self.s.orders()[0]['data']['tracking'], REMOTE['tracking'])

    def test_requery_keeps_confirmed_invoice_and_reports_difference(self):
        job = self.job(); self.query(job)
        self.s.confirm_tracking(job, [self.s.tracking_rows(job)[0]['id']])
        self.query(job, [{**REMOTE, 'tracking': '999999999999'}])
        row = self.s.tracking_rows(job)[0]
        self.assertEqual((row['state'], row['tracking']), ('confirmed', REMOTE['tracking']))
        self.assertIn('확정 송장 유지', row['reason'])
        self.assertEqual(self.s.db.execute('SELECT COUNT(*) FROM tracking_queries').fetchone()[0], 2)

    def test_background_query_detects_concurrent_change(self):
        job = self.job(); expected = encode(self.s.tracking_rows(job))
        row = self.s.tracking_rows(job)[0]
        self.s.confirm_tracking(job, [row['id']], REMOTE['tracking'], '작업자 확인')
        with self.assertRaisesRegex(ValueError, '변경'):
            self.s.apply_tracking_query(job, [REMOTE], DAY, DAY, expected)
        self.assertEqual(self.s.db.execute('SELECT COUNT(*) FROM tracking_queries').fetchone()[0], 0)

    def test_invalid_confirm_rolls_back_and_manual_is_audited(self):
        job = self.job(); row = self.s.tracking_rows(job)[0]
        with self.assertRaises(ValueError): self.s.confirm_tracking(job, [row['id']])
        with self.assertRaises(ValueError): self.s.confirm_tracking(job, [row['id']], REMOTE['tracking'])
        self.s.confirm_tracking(job, [row['id']], REMOTE['tracking'], '위킵 주문정보 대조')
        self.assertTrue(self.s.tracking_rows(job)[0]['confirmed_by'])
        with self.assertRaises(ValueError):
            self.s.confirm_tracking(job, [row['id']], '999999999999', '변경')

    def test_export_preserves_shipping_layout_and_text_identifiers(self):
        job = self.job(); self.query(job)
        path = self.folder / 'tracking.xlsx'
        self.s.export_tracking(job, path)
        book = load_workbook(path)
        self.assertIsNone(book.worksheets[0]['J2'].value)
        book.close()
        self.s.confirm_tracking(job, [self.s.tracking_rows(job)[0]['id']])
        self.s.export_tracking(job, path, confirmed_only=True)
        book = load_workbook(path)
        sheet = book.worksheets[0]
        self.assertEqual([c.value for c in sheet[1]], HEADERS)
        self.assertEqual((sheet['A2'].value, sheet['B2'].value, sheet['E2'].value, sheet['J2'].value),
                         (ORDER['order_no'], ORDER['channel'], ORDER['recipient'], REMOTE['tracking']))
        self.assertEqual(sheet['G2'].value, '00123')
        self.assertEqual(sheet['J2'].data_type, 's')
        self.assertEqual(sheet['J2'].number_format, '@')
        self.assertEqual(book.sheetnames, ['택배출고', '송장 확인 상태'])
        self.assertEqual(sheet.freeze_panes, 'A2')
        book.close()

    def test_original_download_and_workspace_roundtrip(self):
        job = self.job(); self.query(job)
        self.s.confirm_tracking(job, [self.s.tracking_rows(job)[0]['id']])
        path = self.folder / 'original-tracked.xlsx'
        self.s.export_tracking_source(job, path)
        book = load_workbook(path)
        self.assertEqual(book.active.cell(2, 12).value, REMOTE['tracking'])
        book.close()
        state = export_workspace(self.s)
        self.assertEqual(state['schema_version'], 8)
        import_workspace(self.s, state)
        self.assertEqual(workspace_digest(export_workspace(self.s)), workspace_digest(state))

    def test_remote_import_and_formula_like_text_is_not_executed(self):
        path = self.folder / 'remote.xlsx'
        path.write_bytes(workbook_bytes(['판매처주문번호', '송장번호', '수령자'], [['O1', '001234567890', '=HYPERLINK("evil")']]))
        self.assertEqual(read_tracking_file(path)[0]['recipient'], '=HYPERLINK("evil")')
        job = self.job()
        row = self.s.tracking_rows(job)[0]
        data = dict(row['data'], product='=HYPERLINK("evil")')
        with self.s.db:
            self.s.db.execute('UPDATE tracking_rows SET data=? WHERE id=?', (encode(data), row['id']))
        self.s.confirm_tracking(job, [row['id']], REMOTE['tracking'], '확인')
        self.s.export_tracking(job, path)
        book = load_workbook(path)
        self.assertEqual(book.active['C2'].data_type, 's')
        book.close()

    def test_daily_download_combines_marketplaces_and_excludes_other_days(self):
        first = self.job()
        self.s.confirm_tracking(first, [self.s.tracking_rows(first)[0]['id']], REMOTE['tracking'], '확인')
        second = self.s.import_tracking_files([self.path], DAY, channel_override='지마켓')[0][0]
        self.s.import_tracking_files([self.path], '2026-10-07', channel_override='오늘의집')
        path = self.folder / 'all-markets.xlsx'
        self.s.export_tracking_day(DAY, path)
        book = load_workbook(path)
        self.assertEqual(book.active.max_row, 3)
        self.assertEqual({book.active['B2'].value, book.active['B3'].value}, {'오늘의집', '지마켓'})
        book.close()
        self.s.export_tracking_day(DAY, path, confirmed_only=True)
        book = load_workbook(path)
        self.assertEqual(book.active.max_row, 2)
        book.close()


if __name__ == '__main__':
    unittest.main()
