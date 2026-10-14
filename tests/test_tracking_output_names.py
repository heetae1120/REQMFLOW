import unittest
from openpyxl import load_workbook
from tests import test_b2c_daily as fixture
from reqm_local.files import workbook_bytes
from reqm_local.tracking_store import encode

class OutputNameTests(unittest.TestCase):
    def setUp(self):
        fixture.DailyB2CTests.setUp(self)
        self.path=self.folder/'colours.xlsx'
        self.product='맥큐 2in1 Qi2 맥세이프 고속 무선충전 오브제 ACONE 애플워치 갤럭시워치'
        self.path.write_bytes(workbook_bytes(fixture.SOURCE_HEADERS, [
            ['SAME','SAME',self.product,option,1,59900,fixture.ORDER['recipient'],fixture.ORDER['phone'],fixture.ORDER['postcode'],fixture.ORDER['address'],0]
            for option in ('미드나잇섀도우','모닝브리즈')]))
        self.job=self.s.import_tracking_files([self.path],fixture.DAY)[0][0]
        self.s.import_files([self.path])
        for order in self.s.orders():
            code='QMC-ACONE-MS' if order['data']['option']=='미드나잇섀도우' else 'QMC-ACONE-MB'
            self.s.set_mapping(order['id'],[dict(code=code,logistics_code=code,name='예전 품목명',quantity=1,unit_amount=0,warehouse='300',customer='DEMO')])
        with self.s.db:
            self.s.db.execute("UPDATE tracking_rows SET state='confirmed',tracking='001234567890' WHERE job_id=?",(self.job,))

    def tearDown(self):fixture.DailyB2CTests.tearDown(self)

    def test_same_order_colours_use_current_names_in_both_final_exports(self):
        before=encode([self.s.orders(),self.s.tracking_rows(self.job)])
        for export in (lambda p:self.s.export_tracking(self.job,p),lambda p:self.s.export_wekeep_jobs([self.job],p,True)):
            path=self.folder/'final.xlsx';export(path)
            book=load_workbook(path)
            self.assertEqual([book.active.cell(r,3).value for r in (2,3)],['ACONE 미드나잇섀도우','ACONE 모닝브리즈'])
            self.assertEqual([book.active.cell(r,4).value for r in (2,3)],['1','1'])
            self.assertEqual([book.active.cell(r,10).value for r in (2,3)],['001234567890']*2)
            book.close()
        self.assertEqual(encode([self.s.orders(),self.s.tracking_rows(self.job)]),before)

    def test_ambiguous_or_missing_match_keeps_option(self):
        order=self.s.orders()[0]
        with self.s.db:
            self.s.db.execute('INSERT INTO orders SELECT ?,?,data,components,state,issue FROM orders WHERE id=?',('duplicate','duplicate',order['id']))
        rows=self.s.tracking_rows(self.job)
        row=next(r for r in rows if r['data']['option']==order['data']['option'])
        self.assertEqual(self.s.tracking_output_name(row),self.product+' · '+order['data']['option'])
        with self.s.db:self.s.db.execute('DELETE FROM orders')
        self.assertTrue(all(r['data']['option'] in self.s.tracking_output_name(r) for r in rows))

    def test_set_names_show_component_quantities(self):
        order=self.s.orders()[0]
        components=[dict(code='QMC-ACONE-MS',name='',quantity=2),dict(code='QMC-ACONE-MB',name='',quantity=1)]
        with self.s.db:self.s.db.execute('UPDATE orders SET components=? WHERE id=?',(encode(components),order['id']))
        row=next(r for r in self.s.tracking_rows(self.job) if r['data']['option']==order['data']['option'])
        self.assertEqual(self.s.tracking_output_name(row),'ACONE 미드나잇섀도우 × 2개 / ACONE 모닝브리즈 × 1개')
