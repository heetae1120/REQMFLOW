import io
import json
import sqlite3
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from reqm_local.files import (
    DEFAULT_SETTINGS, RESULT_COLUMNS, WEKEEP_REQUEST_COLUMNS, analyze_order_columns,
    choice_column_index, column_choice, filename_signature, parse_orders,
    profile_column_choices, related_column_choices, sample_header_names,
    settings_at, workbook_bytes,
)
from reqm_local.service import Operations
from reqm_local.shipping import ShippingCatalog, compact
from reqm_local.ui_helpers import filter_combobox_choices, search_suggestions
from reqm_local.workspace_cloud import export_workspace, import_workspace, workspace_digest
from reqm_local.updater import RELEASE_API, prepare_update, version_tuple
from reqm_local.desktop import FIELD_LABELS, MAPPING_FIELD_ORDER

REFERENCE=Path(__file__).resolve().parents[1]/'supabase/ecount_migration/data'
HEADERS=['주문번호','상품주문번호','상품명','옵션정보','수량','최종 상품별 총 주문금액','수취인명','수취인연락처1','우편번호','통합배송지','배송비 묶음번호','배송비 합계','주문상태','결제일']


class LocalTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.folder=Path(self.temp.name)
        self.s=Operations(self.folder/'data',REFERENCE)

    def tearDown(self):
        self.s.close();self.temp.cleanup()

    def test_related_search_suggestions_include_direct_and_synonym_terms(self):
        suggestions=search_suggestions('전화',['010-1234-5678','서울 배송지','모트모트 상품'])
        self.assertIn('휴대폰',suggestions)
        self.assertIn('연락처',suggestions)
        direct=search_suggestions('모트',['모트모트 상품','다른 상품'])
        self.assertEqual(direct[0],'모트모트 상품')

    def test_combobox_search_matches_every_typed_term(self):
        choices=['QS-QST01M_BK · 미니셀카봉 블랙','QS-QST01M_WH · 미니셀카봉 화이트','OTHER · 충전기']
        self.assertEqual(filter_combobox_choices('qst 화이트',choices),[choices[1]])

    def test_customer_code_is_remembered_for_the_same_channel(self):
        _,order=self.import_order()
        component=dict(code='TEST',logistics_code='WARE-TEST',name='테스트',quantity=1,unit_amount='0',warehouse='300',customer='SAVED-C')
        self.s.set_mapping(order['id'],[component])
        self.assertEqual(self.s.channel_customer_code(order['data']['channel']),'SAVED-C')
        component['customer']=''
        self.s.set_mapping(order['id'],[component])
        self.assertEqual(self.s.orders()[0]['state'],'출고 준비')

    def test_delivery_values_are_normalized_and_changes_are_recorded(self):
        path=self.folder/'delivery.xlsx'
        path.write_bytes(workbook_bytes(HEADERS,[['O2','L2','테스트상품','기본',1,1000,'수령인','82-010-1234-5678','12345','  서울  테스트로 1  ','B2',0,'결제완료','2026-09-30']]))
        self.s.import_files([path])
        data=self.s.orders()[0]['data']
        self.assertEqual(data['phone'],'01012345678')
        self.assertEqual(data['address'],'서울 테스트로 1')
        self.assertIn('phone',data['delivery_changes'])

    def test_preflight_reports_duplicate_and_match_summary(self):
        _,order=self.import_order()
        self.map(order)
        _,second=self.import_order(line='A2')
        self.map(second)
        # Force the same business row while retaining a distinct database identity.
        with self.s.db:
            self.s.db.execute('UPDATE orders SET data=? WHERE id=?',(json.dumps(order['data'],ensure_ascii=False),second['id']))
        orders=self.s.orders();stats=self.s.match_statistics(orders)
        self.assertEqual(stats['ready'],2)
        issues=self.s.preflight([row['id'] for row in orders])
        self.assertTrue(any(kind=='중복' for _,kind,_ in issues))

    def test_distinct_line_numbers_are_not_duplicate_orders(self):
        _,first=self.import_order(line='A1')
        _,second=self.import_order(line='A2')
        counts={order['data']['line_no']:order['duplicate_count'] for order in self.s.orders()}
        self.assertEqual(counts,{'A1':1,'A2':1})

    def test_force_shipping_approval_allows_review_order_and_records_reason(self):
        _,order=self.import_order()
        self.map(order)
        with self.s.db:
            self.s.db.execute("UPDATE orders SET state='검토 필요',issue='특수 조건 확인' WHERE id=?",(order['id'],))
        self.assertTrue(self.s.preflight([order['id']]))
        self.s.set_force_shipping_approval([order['id']],True,'정상 재구매')
        approved=next(row for row in self.s.orders() if row['id']==order['id'])
        self.assertTrue(approved['data']['force_shipping_approved'])
        self.assertEqual(approved['data']['force_shipping_reason'],'정상 재구매')
        self.assertEqual(self.s.preflight([order['id']]),[])
        self.s.request([order['id']],'2026-10-01',self.folder/'forced.xlsx')
        self.assertTrue((self.folder/'forced.xlsx').exists())

    def test_force_shipping_approval_can_be_revoked(self):
        _,order=self.import_order();self.map(order)
        self.s.set_force_shipping_approval([order['id']],True,'교환 재출고')
        self.s.set_force_shipping_approval([order['id']],False)
        data=self.s.orders()[0]['data']
        self.assertNotIn('force_shipping_approved',data)

    def test_order_can_be_deleted_after_request_and_shipment(self):
        _,order=self.import_order()
        line=self.request(order)
        self.result(line,3)
        self.assertEqual(self.s.orders()[0]['state'],'출고 완료')
        self.assertEqual(self.s.delete_orders([order['id']]),1)
        self.assertEqual(self.s.orders(),[])
        self.assertEqual(self.s.db.execute('SELECT COUNT(*) FROM request_lines').fetchone()[0],0)
        self.assertEqual(self.s.db.execute('SELECT COUNT(*) FROM shipments').fetchone()[0],0)

    def test_updater_compares_semantic_versions(self):
        self.assertGreater(version_tuple('v1.10.0'),version_tuple('1.9.9'))
        self.assertEqual(version_tuple('1.4'),(1,4,0))
        self.assertEqual(RELEASE_API,'https://api.github.com/repos/heetae1120/REQMFLOW/releases/latest')

    def test_updater_verifies_and_extracts_release_package(self):
        import hashlib,zipfile
        package=self.folder/'update.zip'
        with zipfile.ZipFile(package,'w') as archive:
            archive.writestr('REQM_FLOW/REQM_FLOW.exe',b'demo executable')
        checksum=self.folder/'update.sha256'
        checksum.write_text(hashlib.sha256(package.read_bytes()).hexdigest()+'  REQM_FLOW_windows.zip',encoding='utf-8')
        source=prepare_update({'package_url':package.as_uri(),'checksum_url':checksum.as_uri()})
        self.assertEqual((source/'REQM_FLOW.exe').read_bytes(),b'demo executable')

    def test_shared_workspace_snapshot_restores_orders_settings_and_binary_history(self):
        _,order=self.import_order();self.request(order)
        self.s.settings['cloud_email']='local@example.com'
        self.s.settings['profiles'][0]['detected_headers']={'0':'상품주문번호','1':'주문번호'}
        state=export_workspace(self.s)
        digest=workspace_digest(state)
        self.s.db.execute('DELETE FROM shipments')
        self.s.db.execute('DELETE FROM request_lines')
        self.s.db.execute('DELETE FROM requests')
        self.s.db.execute('DELETE FROM artifacts')
        self.s.db.execute('DELETE FROM orders')
        self.s.db.commit()
        self.s.settings['cloud_email']='another@example.com'
        import_workspace(self.s,state)
        self.assertEqual(len(self.s.orders()),1)
        self.assertEqual(len(self.s.artifacts()),1)
        self.assertEqual(self.s.settings['cloud_email'],'another@example.com')
        self.assertEqual(self.s.settings['profiles'][0]['detected_headers']['1'],'주문번호')
        self.assertEqual(workspace_digest(export_workspace(self.s)),digest)

    def import_order(self, q=3, amount=10001, line='A1', fee=3000, bundle='B1', product='테스트상품'):
        path=self.folder/f'{line}.xlsx'
        path.write_bytes(workbook_bytes(HEADERS,[['O1',line,product,'기본',q,amount,'가상수령인','01000000000','00123','테스트 주소',bundle,fee,'결제완료','2026-09-30']]))
        self.s.import_files([path])
        return path,next(o for o in self.s.orders() if o['data']['line_no']==line)

    def map(self, order, components=None):
        self.s.set_mapping(order['id'],components or [dict(code='TEST',logistics_code='WARE-TEST',name='테스트',quantity=1,unit_amount='0',warehouse='300',customer='TEST-C')])

    def request(self, order):
        self.map(order)
        self.s.request([order['id']],'2026-10-01',self.folder/'request.xlsx')
        return self.s.db.execute('SELECT id FROM request_lines WHERE order_id=?',(order['id'],)).fetchone()[0]

    def result(self, line, q, tracking='T1', on='2026-10-01'):
        path=self.folder/'result.xlsx';path.write_bytes(workbook_bytes(RESULT_COLUMNS,[[line,q,tracking,on]]))
        return self.s.import_results(path)

    def test_partial_shipping_erp_exact_total_and_no_duplicates(self):
        _,o=self.import_order();line=self.request(o)
        self.result(line,1)
        self.assertEqual(self.s.orders()[0]['state'],'부분 출고')
        self.s.export_erp('2026-10-01','2026-10-01',self.folder/'first.xlsx')
        self.assertEqual(self.result(line,1),(0,1))
        with self.assertRaises(ValueError):self.s.export_erp('2026-10-01','2026-10-01',self.folder/'duplicate.xlsx')
        self.result(line,2,'T2','2026-10-02')
        batch=self.s.export_erp('2026-10-02','2026-10-02',self.folder/'second.xlsx')
        total=0
        for name in ('first.xlsx','second.xlsx'):
            book=load_workbook(self.folder/name,data_only=True)
            total+=sum(r[15]*r[16] for r in list(book.active.values)[1:]);book.close()
        self.assertEqual(total,13001)
        self.s.mark_registered(batch);self.s.reexport(batch,self.folder/'again.xlsx')
        self.assertEqual((self.folder/'again.xlsx').read_bytes(),(self.folder/'second.xlsx').read_bytes())

    def test_actual_shipment_amount_can_be_rematched_before_erp_export(self):
        _,order=self.import_order(amount=10000,fee=0)
        line=self.request(order)
        self.result(line,3,on='2026-10-05')
        pending=self.s.pending_erp_shipments('2026-10-05')
        self.assertEqual(len(pending),1)
        self.assertEqual(pending[0]['amount'],'10000')
        self.s.set_erp_amount(pending[0]['id'],'12345')
        path=self.folder/'rematched.xlsx'
        self.s.export_erp('2026-10-05','2026-10-05',path)
        book=load_workbook(path,data_only=True)
        total=sum(row[15]*row[16] for row in list(book.active.values)[1:])
        book.close()
        self.assertEqual(total,12345)

    def test_erp_export_matches_ecount_template_columns_and_warehouse_order(self):
        _,warehouse_300=self.import_order(line='A1',fee=0)
        self.map(warehouse_300,[dict(code='ITEM-300',logistics_code='W300',name='위킵품목',quantity=1,unit_amount=0,warehouse='300 위킵창고',customer='C300')])
        self.s.request([warehouse_300['id']],'2026-10-05',self.folder/'request-300.xlsx')
        line_300=self.s.db.execute('SELECT id FROM request_lines WHERE order_id=?',(warehouse_300['id'],)).fetchone()[0]
        self.result(line_300,3,tracking='T300',on='2026-10-05')
        _,warehouse_100=self.import_order(line='A2',fee=0)
        self.map(warehouse_100,[dict(code='ITEM-100',logistics_code='W100',name='본사품목',quantity=1,unit_amount=0,warehouse='100 본사창고',customer='C100')])
        self.s.request([warehouse_100['id']],'2026-10-05',self.folder/'request-100.xlsx')
        line_100=self.s.db.execute('SELECT id FROM request_lines WHERE order_id=?',(warehouse_100['id'],)).fetchone()[0]
        self.result(line_100,3,tracking='T100',on='2026-10-05')
        path=self.folder/'ecount-template.xlsx'
        self.s.export_erp('2026-10-05','2026-10-05',path)
        book=load_workbook(path,data_only=True);rows=list(book['이카운트 웹입력'].values)[1:];book.close()
        warehouses=[row[5] for row in rows]
        self.assertEqual(set(warehouses),{100,300})
        self.assertEqual(warehouses,sorted(warehouses))
        allowed={2,4,5,12,15,16,18,19}
        for row in rows:
            self.assertEqual({index for index,value in enumerate(row) if value not in (None,'')},allowed)
            self.assertIsNone(row[0])
            self.assertIsNone(row[13])

    def test_mapping_rejects_unknown_warehouse(self):
        _,order=self.import_order()
        with self.assertRaisesRegex(ValueError,'100 본사창고 또는 300 위킵창고'):
            self.map(order,[dict(code='ITEM',logistics_code='W',name='품목',quantity=1,unit_amount=0,warehouse='기타창고',customer='C')])

    def test_pending_erp_shipments_are_filtered_by_actual_shipping_day(self):
        _,first=self.import_order(line='A1');line_one=self.request(first)
        self.result(line_one,3,tracking='T1',on='2026-10-04')
        _,second=self.import_order(line='A2');line_two=self.request(second)
        self.result(line_two,3,tracking='T2',on='2026-10-05')
        self.assertEqual([row['order_no'] for row in self.s.pending_erp_shipments('2026-10-05')],['O1'])
        self.assertEqual(len(self.s.pending_erp_shipments()),2)

    def test_duplicate_and_changed_orders_are_allowed_in_test_mode(self):
        path,o=self.import_order()
        self.assertEqual(self.s.import_files([path]),(1,0))
        duplicates=self.s.orders()
        self.assertEqual([order['duplicate_count'] for order in duplicates[:2]],[2,2])
        book=load_workbook(path);book.active['F2']=20000;book.save(path);book.close()
        self.assertEqual(self.s.import_files([path]),(1,0))
        self.assertEqual(len(self.s.orders()),3)

    def test_no_request_for_unmapped_order(self):
        _,o=self.import_order()
        with self.assertRaises(ValueError):self.s.request([o['id']],'2026-10-01',self.folder/'bad.xlsx')
        self.assertFalse((self.folder/'bad.xlsx').exists())

    def test_request_uses_wekeep_layout(self):
        _,order=self.import_order();self.request(order)
        book=load_workbook(self.folder/'request.xlsx',data_only=True)
        sheet=book['택배출고']
        rows=list(sheet.values)
        self.assertEqual(list(rows[0]),WEKEEP_REQUEST_COLUMNS)
        self.assertEqual(rows[1][0:9],('O1','스마트스토어','테스트',3,'가상수령인','01000000000','00123','테스트 주소',None))
        self.assertEqual(rows[1][9],None)
        self.assertEqual(rows[0][10],'일련번호')
        self.assertIsNone(rows[1][10])
        self.assertEqual(sheet.freeze_panes,'A2')
        self.assertEqual(sheet.column_dimensions['C'].width,45)
        book.close()

    def test_first_pass_column_analysis_and_letter_choices(self):
        headers=['주문번호','상품주문번호','상품명','수량','결제금액','수령자','연락처','우편번호','주소']
        sample=[['O1','L1','테스트',2,10000,'홍길동','010-1234-5678','01234','서울시 테스트구']]
        mapping=analyze_order_columns(headers,sample)
        self.assertEqual(mapping['order_no'],0)
        self.assertEqual(mapping['amount'],4)
        self.assertEqual(mapping['phone'],6)
        self.assertEqual(column_choice(26,'추가열'),'AA · 추가열')
        self.assertEqual(choice_column_index('AA · 추가열'),26)
        choices=['미사용','A · 주문번호','B · 휴대폰','C · 배송주소','D · 상품명']
        self.assertEqual(related_column_choices('phone',choices,'전화'),['미사용','B · 휴대폰'])
        self.assertEqual(related_column_choices('address',choices,'배송지'),['미사용','C · 배송주소'])
        self.assertEqual(related_column_choices('order_no',choices,'주문'),['미사용','A · 주문번호'])

    def test_profile_choices_restore_header_labels_for_saved_indexes(self):
        customized = {'column_indexes': {'amount': 28}, 'columns': {}}
        preset = {'columns': {'amount': ['결제금액']}}
        choices = profile_column_choices(customized, preset)
        self.assertEqual(choices, ['미사용', 'AC · 결제금액'])

    def test_profile_choices_keep_all_detected_headers_and_drop_blank_columns(self):
        profile = {
            'detected_headers': {'0': '상품주문번호', '1': '주문번호', '3': '주문상태'},
            'column_indexes': {'order_no': 1},
            'columns': {'order_no': ['주문번호']},
        }
        self.assertEqual(profile_column_choices(profile), [
            '미사용', 'A · 상품주문번호', 'B · 주문번호', 'D · 주문상태',
        ])

    def test_sample_headers_fill_sparse_cells_from_profile(self):
        rows = [['주문번호', '', '수령인'], ['A-1', '10000', '홍길동']]
        profile = {'column_indexes': {'amount': 1}, 'columns': {}}
        preset = {'columns': {'amount': ['결제금액']}}
        self.assertEqual(sample_header_names(rows, 0, profile, preset), ['주문번호', '결제금액', '수령인'])

    def test_no_required_columns_and_missing_ids_get_test_defaults(self):
        profile={**DEFAULT_SETTINGS['profiles'][0], 'columns':{'product':['상품명']}, 'required':[], 'header_row':1}
        path=self.folder/'minimum.xlsx'
        path.write_bytes(workbook_bytes(['상품명'],[['최소 주문']]))
        self.s.settings['profiles']=[profile]
        self.assertEqual(self.s.import_files([path]),(1,0))
        order=self.s.orders()[0]['data']
        self.assertTrue(order['order_no'].startswith('AUTO-ORDER-'))
        self.assertTrue(order['line_no'].startswith('AUTO-LINE-'))
        self.assertEqual(order['quantity'],'1')
        self.assertEqual(order['amount'],'0')

    def test_split_shipping_address_is_joined_with_one_space(self):
        profile = {
            'name':'쌤몰', 'channel':'쌤몰', 'filename_hints':['readyDeliveryList'],
            'required':[], 'header_row':1, 'enabled':True,
            'columns':{
                'order_no':['배송번호'], 'line_no':['품목별 주문번호'],
                'product':['상품명'], 'quantity':['수량'], 'recipient':['수령인명'],
                'phone':['휴대폰'], 'postcode':['우편번호'],
            },
            # Simulate an old customized profile that lost its combine rule.
            'combine':{},
        }
        path=self.folder/'readyDeliveryList.csv'
        path.write_text(
            '배송번호,품목별 주문번호,상품명,수량,수령인명,휴대폰,우편번호,수령인 주소1,수령인 주소2\n'
            'O1,L1,테스트상품,1,홍길동,01000000000,12345," 강원 원주시  가곡로 50 ","1006동 703호  (원주롯데캐슬더퍼스트)"\n',
            encoding='utf-8-sig',
        )
        self.s.settings['profiles']=[profile]
        self.assertEqual(self.s.import_files([path]),(1,0))
        order=self.s.orders()[0]
        self.assertEqual(order['data']['address'],'강원 원주시 가곡로 50 1006동 703호 (원주롯데캐슬더퍼스트)')
        self.map(order)
        self.s.request([order['id']],'2026-10-02',self.folder/'wekeep-address.xlsx')
        book=load_workbook(self.folder/'wekeep-address.xlsx',data_only=True)
        self.assertEqual(book['택배출고']['H2'].value,'강원 원주시 가곡로 50 1006동 703호 (원주롯데캐슬더퍼스트)')
        book.close()

    def test_mapping_fields_have_fixed_requested_order(self):
        self.assertEqual([FIELD_LABELS[field] for field in MAPPING_FIELD_ORDER], [
            '주문일자', '판매처주문번호', '상품주문번호', '상품명', '옵션', '수량', '금액', '배송비',
            '수령인', '전화번호', '우편번호', '주소 1', '주소 2', '배송메모',
        ])

    def test_explicit_address_columns_are_joined_with_one_space(self):
        profile = {
            'name':'테스트몰', 'channel':'테스트몰', 'filename_hints':['address-parts'],
            'required':[], 'header_row':1, 'enabled':True,
            'columns':{
                'order_no':['주문번호'], 'line_no':['상품주문번호'], 'product':['상품명'],
                'quantity':['수량'], 'address1':['주소 앞부분'], 'address2':['상세 위치'],
            },
        }
        path=self.folder/'address-parts.xlsx'
        path.write_bytes(workbook_bytes(
            ['주문번호','상품주문번호','상품명','수량','주소 앞부분','상세 위치'],
            [['O1','L1','테스트상품',1,' 강원 원주시  가곡로 50 ','1006동 703호  ']],
        ))
        parsed=parse_orders(path,[profile])[0]
        self.assertEqual(parsed['address'],'강원 원주시 가곡로 50 1006동 703호')

    def test_existing_settings_disable_required_fields(self):
        folder=self.folder/'settings-only';folder.mkdir()
        (folder/'settings.json').write_text(json.dumps(DEFAULT_SETTINGS,ensure_ascii=False),encoding='utf-8')
        settings=json.loads((folder/'settings.json').read_text(encoding='utf-8'))
        settings['profiles'][0]['required']=['주문번호']
        (folder/'settings.json').write_text(json.dumps(settings,ensure_ascii=False),encoding='utf-8')
        self.assertEqual(settings_at(folder)['profiles'][0]['required'],[])

    def test_delete_pending_and_requested_orders(self):
        _,first=self.import_order(line='DELETE-1')
        self.assertEqual(self.s.delete_orders([first['id']]),1)
        self.assertEqual(len(self.s.orders()),0)
        _,second=self.import_order(line='DELETE-2');self.request(second)
        self.assertEqual(self.s.delete_orders([second['id']]),1)
        self.assertEqual(self.s.orders(),[])

    def test_event_rule_transforms_matching_pending_and_new_orders(self):
        path,order=self.import_order(line='EVENT-1',amount=12000,product='행사 전 상품')
        components=[dict(code='EVENT-A',logistics_code='W-EVENT-A',name='이벤트 본품',quantity=1,unit_amount=0,warehouse='300',customer='EVENT-C'),
                    dict(code='EVENT-GIFT',logistics_code='W-GIFT',name='이벤트 사은품',quantity=2,unit_amount=500,warehouse='300',customer='EVENT-C')]
        rule=self.s.set_event_rule(order['id'],'가을 행사','행사 상품','특가 옵션','7900',components)
        changed_order=self.s.orders()[0];changed=changed_order['data']
        self.assertEqual((changed['event_name'],changed['product'],changed['option'],changed['amount']),('가을 행사','행사 상품','특가 옵션','23700'))
        self.assertEqual([(item['code'],item['quantity'],item['amount']) for item in changed_order['components']],
                         [('EVENT-A',3,'22200'),('EVENT-GIFT',3,'1500')])
        second=self.folder/'event-second.xlsx'
        second.write_bytes(workbook_bytes(HEADERS,[['O2','EVENT-2','행사 전 상품','기본',1,12000,'가상수령인','01000000000','00123','테스트 주소','B2',0,'결제완료','2026-10-01']]))
        self.s.import_files([second])
        imported=next(item for item in self.s.orders() if item['data']['line_no']=='EVENT-2')['data']
        self.assertEqual(imported['event_rule_id'],rule)
        self.assertEqual(imported['amount'],'7900')
        imported_order=next(item for item in self.s.orders() if item['data']['line_no']=='EVENT-2')
        self.assertEqual([item['code'] for item in imported_order['components']],['EVENT-A','EVENT-GIFT'])
        self.s.clear_event_rule(order['id'])
        reverted=next(item for item in self.s.orders() if item['data']['line_no']=='EVENT-1')['data']
        self.assertEqual((reverted['product'],reverted['option'],reverted['amount']),('행사 전 상품','기본','12000'))
        self.assertNotIn('event_name',reverted)

    def test_old_event_rule_schema_adds_component_storage(self):
        folder=self.folder/'old-event-data';folder.mkdir()
        db=sqlite3.connect(folder/'operations.sqlite3')
        db.execute('''CREATE TABLE event_rules(id TEXT PRIMARY KEY, channel TEXT NOT NULL, name TEXT NOT NULL,
            source_product TEXT NOT NULL, source_option TEXT NOT NULL, target_product TEXT NOT NULL,
            target_option TEXT NOT NULL, target_amount TEXT NOT NULL, active INTEGER NOT NULL DEFAULT 1,
            UNIQUE(channel, source_product, source_option))''')
        db.commit();db.close()
        service=Operations(folder,REFERENCE)
        try:
            columns={row['name'] for row in service.db.execute('PRAGMA table_info(event_rules)')}
            self.assertIn('components',columns)
        finally:
            service.close()

    def test_over_shipment_rolls_back_whole_file(self):
        _,o=self.import_order();line=self.request(o)
        path=self.folder/'bad.xlsx';path.write_bytes(workbook_bytes(RESULT_COLUMNS,[[line,1,'T1','2026-10-01'],[line,3,'T2','2026-10-01']]))
        with self.assertRaises(ValueError):self.s.import_results(path)
        self.assertEqual(self.s.db.execute('SELECT COUNT(*) FROM shipments').fetchone()[0],0)

    def test_set_quantities_and_price_allocation(self):
        _,o=self.import_order(q=2,amount=20000,fee=0)
        self.map(o,[dict(code='MAIN',logistics_code='M',name='본품',quantity=1,unit_amount=0,warehouse='300',customer='C'),dict(code='CASE',logistics_code='C',name='케이스',quantity=2,unit_amount=1000,warehouse='300',customer='C')])
        result=self.s.orders()[0]['components']
        self.assertEqual([c['quantity'] for c in result],[2,2])
        self.assertEqual([c['amount'] for c in result],['18000','2000'])

    def test_shipping_alias_requires_verified_skus_and_keeps_source_quantity(self):
        source='배터리 케이블 기본'
        self.s.catalog.shipping_catalog=ShippingCatalog(
            items=[{'item_code':'BATTERY','standard_name':'배터리','is_active':True},
                   {'item_code':'CABLE','standard_name':'케이블','is_active':True}],
            aliases=[{'source_channel':'리큐엠_스마트스토어','normalized_source':compact(source),
                      'components':[{'item_code':'BATTERY','quantity':1},{'item_code':'CABLE','quantity':2}],
                      'is_active':True}],
            sku_mappings=[{'item_code':'BATTERY','product_name':'위킵 배터리','sku_no':'SKU-B','is_active':True},
                          {'item_code':'CABLE','product_name':'위킵 케이블','sku_no':'SKU-C','is_active':True}],
        )
        _,order=self.import_order(q=2,amount=20000,product='배터리 케이블')
        ready=self.s.orders()[0]
        self.assertEqual(ready['state'],'출고 준비')
        self.assertEqual([(row['code'],row['quantity'],row['sku_no']) for row in ready['components']],
                         [('BATTERY',2,'SKU-B'),('CABLE',2,'SKU-C')])
        self.s.request([order['id']],'2026-10-01',self.folder/'safe-request.xlsx')
        book=load_workbook(self.folder/'safe-request.xlsx',data_only=True)
        rows=list(book['택배출고'].values)
        self.assertEqual([(row[2],row[3]) for row in rows[1:]], [('위킵 배터리',2),('위킵 케이블',2)])
        book.close()

    def test_shipping_alias_with_missing_sku_is_never_exportable(self):
        self.s.catalog.shipping_catalog=ShippingCatalog(
            items=[{'item_code':'BATTERY','standard_name':'배터리','is_active':True},
                   {'item_code':'CABLE','standard_name':'케이블','is_active':True}],
            aliases=[{'source_channel':'리큐엠_스마트스토어','normalized_source':compact('배터리 케이블 기본'),
                      'components':[{'item_code':'BATTERY'},{'item_code':'CABLE'}],'is_active':True}],
            sku_mappings=[{'item_code':'BATTERY','product_name':'위킵 배터리','sku_no':'SKU-B','is_active':True}],
        )
        _,order=self.import_order(q=2,product='배터리 케이블')
        blocked=self.s.orders()[0]
        self.assertEqual(blocked['state'],'검토 필요')
        self.assertIn('위킵 SKU 미등록 CABLE',blocked['issue'])
        with self.assertRaises(ValueError):
            self.s.request([order['id']],'2026-10-01',self.folder/'must-not-exist.xlsx')
        self.assertFalse((self.folder/'must-not-exist.xlsx').exists())

    def test_incomplete_shipping_catalog_fails_closed(self):
        self.s.catalog.shipping_catalog=ShippingCatalog(load_errors=['item_aliases'])
        self.import_order(product='알 수 없는 상품')
        blocked=self.s.orders()[0]
        self.assertEqual(blocked['state'],'검토 필요')
        self.assertIn('출고 DB를 모두 읽지 못했습니다',blocked['issue'])

    def test_request_rejects_stale_multiplied_component_quantity(self):
        _,order=self.import_order(q=2)
        self.map(order)
        current=self.s.orders()[0]
        stale=[{**current['components'][0],'quantity':4}]
        self.s.db.execute('UPDATE orders SET components=? WHERE id=?',(json.dumps(stale,ensure_ascii=False),order['id']))
        self.s.db.commit()
        with self.assertRaisesRegex(ValueError,'원본 엑셀 수량과 다릅니다'):
            self.s.request([order['id']],'2026-10-01',self.folder/'stale.xlsx')
        self.assertFalse((self.folder/'stale.xlsx').exists())

    def test_request_repetition_and_invalid_date(self):
        _,o=self.import_order();self.request(o)
        with self.assertRaises(ValueError):self.s.request([o['id']],'2026-10-01',self.folder/'again.xlsx')
        with self.assertRaises(ValueError):self.s.export_erp('invalid','2026-10-01',self.folder/'bad.xlsx')

    def test_formula_injection_and_leading_zero_preserved(self):
        content=workbook_bytes(['주소','전화'],[['=HYPERLINK("x")','00123']])
        book=load_workbook(io.BytesIO(content),data_only=False)
        self.assertEqual(book.active['A2'].data_type,'s');self.assertEqual(book.active['B2'].value,'00123');book.close()

    def test_account_identity_and_ambiguous_detection(self):
        path,o=self.import_order()
        profile={**DEFAULT_SETTINGS['profiles'][0],'account':'두번째'}
        with self.assertRaises(ValueError):parse_orders(path,[DEFAULT_SETTINGS['profiles'][0],profile])
        self.s.settings['profiles']=[profile]
        self.assertEqual(self.s.import_files([path]),(1,0))

    def test_manual_channel_import_uses_selected_channel(self):
        path=self.folder/'today-house.xlsx'
        path.write_bytes(workbook_bytes(HEADERS,[['O2','T1','테스트상품','기본',1,12000,'가상수령인','01000000000','00123','테스트 주소','B2',0,'결제완료','2026-10-01']]))
        self.assertEqual(self.s.import_files([path],channel_override='오늘의집'),(1,0))
        self.assertEqual(self.s.orders()[0]['data']['channel'],'오늘의집')

    def test_filename_matching_ignores_changing_date(self):
        path=self.folder/'오늘의집_주문_20261001.xlsx'
        path.write_bytes(workbook_bytes(HEADERS,[['O3','T2','테스트상품','기본',1,12000,'가상수령인','01000000000','00123','테스트 주소','B3',0,'결제완료','2026-10-01']]))
        base=DEFAULT_SETTINGS['profiles'][0]
        profiles=[
            {**base,'name':'오늘의집','channel':'오늘의집','filename_hints':['오늘의집_주문_20260930.xlsx']},
            {**base,'name':'에이블리','channel':'에이블리','filename_hints':['에이블리_주문_20260930.xlsx']},
        ]
        self.assertEqual(parse_orders(path,profiles)[0]['channel'],'오늘의집')
        self.assertEqual(filename_signature('11번가_주문_20261001.xlsx'),'11번가 주문')
        self.assertEqual(filename_signature('29CM-orders-2026-10-01.xlsx'),'29cm orders')

    def test_smartstore_n_delivery_rows_are_excluded(self):
        headers=HEADERS+['배송속성']
        path=self.folder/'스마트스토어_전체주문배송현황_20261003.xlsx'
        common=['O1','L1','상품','기본',1,10000,'수령인','01000000000','00123','서울','B1',0,'결제완료','2026-10-03']
        path.write_bytes(workbook_bytes(headers,[common+['N배송'],[{**dict(enumerate(common)),1:'L2'}.get(i,'') for i in range(len(common))]+['일반배송']]))
        rows=parse_orders(path,[DEFAULT_SETTINGS['profiles'][0]])
        self.assertEqual([row['line_no'] for row in rows],['L2'])

    def test_identical_samsung_filenames_use_ba_order_prefix(self):
        path=self.folder/'Excel_20261003120000.xlsx'
        path.write_bytes(workbook_bytes(['상품명','수량','주문번호'],[['상품',1,'BA12345']]))
        base={
            'enabled':True,'filename_hints':['Excel_20260907151205.xlsx'],'required':[],
            'columns':{'order_no':['주문번호'],'product':['상품명'],'quantity':['수량']},
        }
        profiles=[
            {**base,'name':'삼성카드복지몰','channel':'삼성카드복지몰','content_rule':{'column_index':2,'prefix':'BA'}},
            {**base,'name':'삼성카드쇼핑몰','channel':'삼성카드쇼핑몰','content_rule':{'not':{'column_index':2,'prefix':'BA'}}},
        ]
        self.assertEqual(parse_orders(path,profiles)[0]['channel'],'삼성카드복지몰')

    def test_persistence_backup(self):
        _,o=self.import_order();self.request(o)
        self.s.backup(self.folder/'backup.sqlite3')
        db=sqlite3.connect(self.folder/'backup.sqlite3')
        self.assertEqual(db.execute('SELECT COUNT(*) FROM request_lines').fetchone()[0],1);db.close()
        self.s.close();self.s=Operations(self.folder/'data',REFERENCE)
        self.assertEqual(self.s.orders()[0]['state'],'출고 요청')

    def test_shared_shipping_fee_only_when_bundle_complete(self):
        _,first=self.import_order(line='A1');_,second=self.import_order(line='A2')
        one=self.request(first);two=self.request(second)
        self.result(one,3);self.s.export_erp('2026-10-01','2026-10-01',self.folder/'first.xlsx')
        self.assertEqual(self.s.db.execute('SELECT COUNT(*) FROM fees').fetchone()[0],0)
        self.result(two,3);self.s.export_erp('2026-10-01','2026-10-01',self.folder/'second.xlsx')
        self.assertEqual(self.s.db.execute('SELECT COUNT(*) FROM fees').fetchone()[0],1)

    def test_cancelled_order_cannot_be_released_by_mapping(self):
        path,o=self.import_order()
        d=o['data'];d['status']='취소완료'
        with self.s.db:self.s.db.execute('UPDATE orders SET data=? WHERE id=?',(json.dumps(d),o['id']))
        self.map(o)
        self.assertEqual(self.s.orders()[0]['state'],'검토 필요')

if __name__=='__main__':unittest.main()
