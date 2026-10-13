import io
import json
import tempfile
import unittest
import uuid
from pathlib import Path
from decimal import Decimal
from openpyxl import load_workbook
from reqm_local.service import Operations, encode
from reqm_local.workspace_cloud import export_workspace, import_workspace, workspace_digest
from tests.test_tracking import ORDER, DAY, REFERENCE


class VoucherTests(unittest.TestCase):
    def test_live_allocation_remaining_equal_excess_and_invalid(self):
        from reqm_local.voucher import allocation_feedback
        self.assertIn('남은 금액 8,400원',allocation_feedback('36,900',['28,500',''])[0])
        self.assertTrue(allocation_feedback('36,900',['28,500','8,400'])[2])
        self.assertIn('초과 금액 100원',allocation_feedback('36,900',['28,500','8,500'])[0])
        for value in ('-1','NaN','Infinity','1.5','문자'):
            with self.assertRaises(ValueError):allocation_feedback(36900,[value])

    def test_price_review_quantity_scaling_reason_and_export(self):
        a=self.order();self.s.set_voucher_amounts(a['id'],34900,[28000,6900])
        scaled=self.order('S',total=69800,qty=2)
        self.assertFalse(self.s.voucher_review(scaled,69800,[56000,13800]))
        changed=self.order('B',total=33900)
        self.assertIn('가격 차이 확인',self.s.voucher_review(changed)[0])
        self.assertIsNone(self.s.voucher_suggestion(changed))
        with self.assertRaisesRegex(ValueError,'확인 사유'):
            self.s.set_voucher_amounts(changed['id'],33900,[27000,6900])
        self.s.set_voucher_amounts(changed['id'],33900,[27000,6900],review_reason='행사 할인 1,000원 확인')
        self.s.set_voucher_amounts(scaled['id'],69800,[56000,13800])
        self.s.export_erp(DAY,DAY,None,from_day=DAY,preview=True)
        updated=next(r for r in self.s.voucher_rows(DAY,DAY) if r['id']==changed['id'])
        self.assertFalse(self.s.voucher_review_pending(updated))
        other=Operations(self.folder/'other',REFERENCE)
        try:
            import_workspace(other,export_workspace(self.s))
            copied=next(r for r in other.voucher_rows(DAY,DAY) if r['id']==changed['id'])
            self.assertFalse(other.voucher_review_pending(copied))
        finally:other.close()

    def test_export_blocks_unreviewed_source_change_without_writes(self):
        a=self.order();self.s.set_voucher_amounts(a['id'],34900,[28000,6900])
        data={**a['data'],'amount':'35900'}
        with self.s.db:self.s.db.execute('UPDATE orders SET data=? WHERE id=?',(encode(data),a['order_id']))
        before=workspace_digest(export_workspace(self.s))
        with self.assertRaisesRegex(ValueError,'원본 금액 불일치'):
            self.s.export_erp(DAY,DAY,self.folder/'blocked.xlsx',from_day=DAY)
        self.assertEqual(workspace_digest(export_workspace(self.s)),before)
        self.assertFalse((self.folder/'blocked.xlsx').exists())

    def test_review_does_not_compare_other_channels_models_partial_or_future(self):
        a=self.order();self.s.set_voucher_amounts(a['id'],34900,[28000,6900])
        b=self.order('B',total=33900,channel='무신사')
        self.assertFalse(self.s.voucher_review(b))
        changed={**a,'id':'new','total':Decimal(33900),'components':[dict(c) for c in a['components']]}
        changed['components'][0]['amount']='27000'
        changed['components'][1]['amount']='6900'
        changed['data']={**a['data'],'amount':'33900'}
        self.assertFalse(self.s.voucher_review({**changed,'full':False}))
        self.assertFalse(self.s.voucher_review(changed,history=[{**a,'day':'2099-01-01'}]))
        changed['components'][0]['code']='[RQM]-QP1000C1-WH'
        self.assertFalse(self.s.voucher_review(changed))

    def test_amount_dialog_updates_remaining_and_save_state(self):
        import os,tkinter as tk
        from tkinter import ttk
        from unittest.mock import patch
        from reqm_local.desktop import Desktop
        a=self.order(total=36900)
        root=tk.Tk();root.withdraw()
        try:
            with patch.dict(os.environ,{'REQM_SELF_TEST':'1'}):app=Desktop(root,self.s)
            panel=app.voucher_panel;panel.start.set(DAY);panel.end.set(DAY);panel.refresh()
            panel.tree.selection_set(panel.tree.get_children()[0]);panel.edit()
            win=next(w for w in root.winfo_children() if isinstance(w,tk.Toplevel))
            frame=win.winfo_children()[0]
            entries=[w for w in frame.winfo_children() if isinstance(w,ttk.Entry)]
            button=next(w for w in frame.winfo_children() if isinstance(w,ttk.Button) and w.cget('text')=='금액 확정')
            def set_entry(widget,value):root.setvar(widget.cget('textvariable'),value)
            def texts():
                return [root.getvar(w.cget('textvariable')) for w in frame.winfo_children() if isinstance(w,ttk.Label) and w.cget('textvariable')]
            set_entry(entries[1],'28,500');set_entry(entries[2],'')
            self.assertTrue(any('남은 금액 8,400원' in text for text in texts()))
            self.assertTrue(button.instate(['disabled']))
            set_entry(entries[2],'8,400')
            self.assertTrue(any('금액 일치' in text for text in texts()))
            self.assertTrue(button.instate(['!disabled']))
            set_entry(entries[2],'8,500')
            self.assertTrue(any('초과 금액 100원' in text for text in texts()))
            self.assertTrue(button.instate(['disabled']))
            set_entry(entries[2],'8,400');set_entry(entries[0],'37,900');set_entry(entries[2],'9,400')
            self.assertTrue(button.instate(['disabled']))
            set_entry(entries[3],'원본 가격 정정 확인')
            self.assertTrue(button.instate(['!disabled']))
            button.invoke()
            updated=next(r for r in self.s.voucher_rows(DAY,DAY) if r['id']==a['id'])
            self.assertEqual(updated['total'],Decimal(37900))
            self.assertFalse(self.s.voucher_review_pending(updated))
        finally:
            for timer in root.tk.call('after','info'):root.after_cancel(timer)
            root.destroy()

    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.folder=Path(self.temp.name)
        self.s=Operations(self.folder/'data', REFERENCE)

    def tearDown(self):
        self.s.close();self.temp.cleanup()

    def order(self, number='A', white=False, total=34900, qty=1, channel='오늘의집', fee=3000, bundle=''):
        oid=uuid.uuid4().hex
        data={**ORDER,'channel':channel,'account':'worker','order_no':number,'quantity':str(qty),
              'amount':str(total),'shipping':str(fee),'bundle':bundle,'memo':''}
        codes=['[RQM]-QP2000C1-WH' if white else '[RQM]-QP2000C1-BK',
               'RQM-CC150-CK-PK' if white else 'RQM-CC150-CK-BL']
        components=[{'code':code,'logistics_code':code,'name':self.s.item_name(code),
                     'quantity':qty,'amount':str(total if i==0 else 0),'warehouse':'300','customer':'DEMO'}
                    for i,code in enumerate(codes)]
        with self.s.db:
            self.s.db.execute('INSERT INTO orders VALUES(?,?,?,?,?,?)',(oid,oid,encode(data),encode(components),'출고 준비',''))
        batch=self.s.request([oid],DAY)
        self.s.auto_import_shipments(DAY)
        return next(r for r in self.s.voucher_rows(DAY,DAY) if r['order_id']==oid)

    def test_other_colour_reuses_allocation_and_scales_order_quantity(self):
        a=self.order();b=self.order('B',True,total=69800,qty=2)
        self.s.set_voucher_amounts(a['id'],34900,[28000,6900])
        self.assertEqual(self.s.reuse_voucher_amounts(DAY,DAY),1)
        b=next(r for r in self.s.voucher_rows(DAY,DAY) if r['id']==b['id'])
        self.assertEqual([c['amount'] for c in b['components']],['56000','13800'])
        self.assertEqual(b['components'][0]['code'],'[RQM]-QP2000C1-WH')
        self.assertTrue(b['confirmed'])

    def test_changed_price_channel_model_and_partial_shipment_do_not_reuse(self):
        a=self.order();self.s.set_voucher_amounts(a['id'],34900,[28000,6900])
        b=self.order('B',True,total=33900);c=self.order('C',True,channel='무신사')
        self.assertIsNone(self.s.voucher_suggestion(b));self.assertIsNone(self.s.voucher_suggestion(c))
        changed={**a,'components':[dict(p) for p in a['components']]}
        changed['components'][0]['code']='[RQM]-QP1000C1-WH'
        self.assertIsNone(self.s.voucher_suggestion(changed))
        self.assertIsNone(self.s.voucher_suggestion({**a,'full':False}))

    def test_conflicting_history_requires_manual_allocation(self):
        a=self.order();b=self.order('B',True);c=self.order('C',True)
        self.s.set_voucher_amounts(a['id'],34900,[28000,6900])
        self.s.set_voucher_amounts(b['id'],34900,[29000,5900],review_reason='배분 기준 변경 확인')
        self.assertIsNone(self.s.voucher_suggestion(c))
        self.assertEqual(self.s.reuse_voucher_amounts(DAY,DAY),0)

    def test_total_validation_is_atomic_and_manual_total_is_supported(self):
        a=self.order();before=workspace_digest(export_workspace(self.s))
        for total,amounts in ((34900,[28000,6000]),(34900,[-1,34901])):
            with self.assertRaises(ValueError):self.s.set_voucher_amounts(a['id'],total,amounts)
            self.assertEqual(workspace_digest(export_workspace(self.s)),before)
        with self.assertRaisesRegex(ValueError,'원본 금액 불일치'):
            self.s.set_voucher_amounts(a['id'],40000,[30000,10000])
        self.assertEqual(workspace_digest(export_workspace(self.s)),before)
        self.s.set_voucher_amounts(a['id'],40000,[30000,10000],review_reason='원본 가격 정정 확인')
        updated=self.s.voucher_rows(DAY,DAY)[0]
        self.assertEqual(updated['total'],Decimal(40000))

    def test_preview_does_not_mark_rows_and_fee_once_per_order_not_bundle(self):
        a=self.order(bundle='B1');b=self.order('A',True,bundle='B2')
        for row in (a,b):self.s.set_voucher_amounts(row['id'],34900,[28000,6900])
        before=workspace_digest(export_workspace(self.s))
        rows=self.s.export_erp(DAY,DAY,None,from_day=DAY,preview=True)
        self.assertEqual(workspace_digest(export_workspace(self.s)),before)
        fees=[r for r in rows if r[12]=='택배운송비']
        self.assertEqual([(r[15],r[16]) for r in fees],[(1,3000)])
        self.assertTrue(all(r[0]==int(DAY.replace('-','')) and r[13] for r in rows))
        target=self.folder/'voucher.xlsx'
        batch=self.s.export_erp(DAY,DAY,target,from_day=DAY)
        book=load_workbook(target);self.assertEqual(book.active['A2'].value,int(DAY.replace('-','')));book.close()
        with self.assertRaises(ValueError):self.s.export_erp(DAY,DAY,target,from_day=DAY)
        again=self.folder/'again.xlsx';self.s.regenerate_erp(batch,again)
        book=load_workbook(again)
        self.assertEqual(sum(r[12]=='택배운송비' for r in book.active.iter_rows(min_row=2,values_only=True)),1)
        book.close()

    def test_shipping_conflict_rolls_back_and_override_is_shared(self):
        a=self.order(fee=3000);b=self.order('A',True,fee=4000)
        for row in (a,b):self.s.set_voucher_amounts(row['id'],34900,[28000,6900])
        before=workspace_digest(export_workspace(self.s))
        with self.assertRaisesRegex(ValueError,'배송비'):self.s.export_erp(DAY,DAY,self.folder/'bad.xlsx',from_day=DAY)
        self.assertEqual(workspace_digest(export_workspace(self.s)),before)
        self.assertFalse((self.folder/'bad.xlsx').exists())
        self.s.set_voucher_fee(a['data'],3500)
        rows=self.s.export_erp(DAY,DAY,None,from_day=DAY,preview=True)
        self.assertEqual([r[16] for r in rows if r[12]=='택배운송비'],[3500])
        other=Operations(self.folder/'other',REFERENCE)
        try:
            import_workspace(other,export_workspace(self.s))
            self.assertEqual(other.voucher_fee(a['data']),Decimal(3500))
        finally:other.close()

    def test_first_partial_export_charges_fee_only_once_and_legacy_fee_is_respected(self):
        a=self.order();self.s.set_voucher_amounts(a['id'],34900,[28000,6900])
        sid=a['components'][1]['entry_id']
        with self.s.db:self.s.db.execute('UPDATE shipments SET day=? WHERE id=?',('2026-10-14',sid))
        self.s.export_erp(DAY,DAY,self.folder/'first.xlsx',from_day=DAY)
        rows=self.s.export_erp('2026-10-14','2026-10-14',None,from_day='2026-10-14',preview=True)
        self.assertFalse(any(r[12]=='택배운송비' for r in rows))

    def test_opt_out_is_not_learned_on_next_batch(self):
        a=self.order();b=self.order('B',True)
        self.s.set_voucher_amounts(a['id'],34900,[28000,6900],remember=False)
        self.assertEqual(self.s.reuse_voucher_amounts(DAY,DAY),0)
        self.assertIsNone(self.s.voucher_suggestion(b))

    def test_smartstore_and_esm_shipping_are_included_without_bundle_column(self):
        for index,table in enumerate(('smartstore_erp_rows','esm_erp_rows')):
            data={**ORDER,'channel':'스마트스토어' if index==0 else '옥션','account':'worker',
                  'order_no':'DIRECT-'+str(index),'amount':'10000','shipping':'3000','quantity':'1'}
            component={'code':'[RQM]-QP2000C1-BK','name':'QP2000C 블랙','quantity':1,
                       'amount':'10000','customer':'DEMO','warehouse':'300'}
            with self.s.db:
                if index==0:
                    self.s.db.execute('INSERT INTO smartstore_erp_rows(id,identity,day,data,components,issue,amount_split_confirmed) VALUES(?,?,?,?,?,?,?)',
                                      ('SS','SS',DAY,encode(data),encode([component]),'',1))
                else:
                    self.s.db.execute('INSERT INTO esm_erp_rows VALUES(?,?,?,?,?,?,?)',
                                      ('ESM','ESM',DAY,encode(data),encode([component]),None,''))
        rows=self.s.export_erp(DAY,DAY,None,from_day=DAY,preview=True)
        self.assertEqual(sum(r[15]*r[16] for r in rows),26000)
        self.assertEqual(sum(r[15] for r in rows if r[12]=='택배운송비'),2)

    def test_rules_round_trip_and_old_confirmed_splits_are_learned(self):
        a=self.order();b=self.order('B',True)
        # Simulate an allocation entered through the pre-existing split screen.
        self.s.set_amount_split(a['order_id'],DAY,{0:28000,1:6900})
        self.assertEqual(self.s.reuse_voucher_amounts(DAY,DAY),1)
        other=Operations(self.folder/'other',REFERENCE)
        try:
            import_workspace(other,export_workspace(self.s))
            self.assertEqual(other.voucher_suggestion(b),['28000','6900'])
        finally:other.close()

    def test_legacy_fee_regeneration_keeps_one_fee_for_a_shared_bundle(self):
        a=self.order(bundle='BUNDLE');b=self.order('B',True,bundle='BUNDLE')
        for row in (a,b):self.s.set_voucher_amounts(row['id'],34900,[28000,6900])
        batch=self.s.export_erp(DAY,DAY,self.folder/'initial.xlsx',from_day=DAY)
        with self.s.db:
            self.s.db.execute('DELETE FROM fees')
            self.s.db.execute('INSERT INTO fees VALUES(?,?)',(self.s.bundle_key(a['data']),batch))
        target=self.folder/'legacy.xlsx';self.s.regenerate_erp(batch,target)
        book=load_workbook(target)
        self.assertEqual(sum(r[15] for r in book.active.iter_rows(min_row=2,values_only=True) if r[12]=='택배운송비'),1)
        book.close()

    def test_new_work_table_and_preview_open_without_modifying_data(self):
        import os
        import tkinter as tk
        from unittest.mock import patch
        from reqm_local.desktop import Desktop
        a=self.order();self.s.set_voucher_amounts(a['id'],34900,[28000,6900])
        root=tk.Tk();root.withdraw()
        try:
            with patch.dict(os.environ,{'REQM_SELF_TEST':'1'}):app=Desktop(root,self.s)
            app.voucher_panel.start.set(DAY);app.voucher_panel.end.set(DAY)
            app.voucher_panel.refresh();root.update_idletasks()
            before=workspace_digest(export_workspace(self.s))
            app.voucher_panel.preview();root.update_idletasks()
            self.assertEqual(len(app.voucher_panel.tree.get_children()),2)
            self.assertEqual(workspace_digest(export_workspace(self.s)),before)
        finally:root.destroy()

    def test_explicit_family_override_and_disable_are_saved_in_workspace(self):
        code='[RQM]-QP2000C1-BK'
        self.assertEqual(self.s.voucher_family(code),self.s.voucher_family('[RQM]-QP2000C1-WH'))
        self.assertNotEqual(self.s.voucher_family(code),self.s.voucher_family('[RQM]-QP1000C1-WH'))
        self.s.set_voucher_family(code,'')
        self.assertEqual(self.s.voucher_family(code),'sku:'+code)
        other=Operations(self.folder/'other',REFERENCE)
        try:
            import_workspace(other,export_workspace(self.s))
            self.assertEqual(other.voucher_family(code),'sku:'+code)
        finally:other.close()

    def test_requests_are_visible_before_confirmation_without_becoming_erp_rows(self):
        a=self.order()
        with self.s.db:
            self.s.db.execute('DELETE FROM shipments')
            self.s.db.execute("UPDATE orders SET state='출고 요청'")
        before=workspace_digest(export_workspace(self.s))
        rows=self.s.voucher_rows(DAY,DAY,include_requests=True)
        self.assertEqual(len(rows),1);self.assertEqual(rows[0]['kind'],'request')
        self.assertEqual(len(rows[0]['components']),2)
        self.assertEqual(workspace_digest(export_workspace(self.s)),before)
        with self.assertRaises(ValueError):self.s.export_erp(DAY,DAY,None,from_day=DAY,preview=True)
        self.assertEqual(self.s.confirm_voucher_requests([rows[0]['id']],'2026-10-09'),2)
        self.assertEqual(self.s.voucher_rows('2026-10-09','2026-10-09')[0]['kind'],'shipment')
        with self.assertRaises(ValueError):self.s.confirm_voucher_requests([rows[0]['id']],'2026-10-09')

    def test_unmapped_smartstore_shows_source_product_and_quantity(self):
        import os,tkinter as tk
        from unittest.mock import patch
        from reqm_local.desktop import Desktop
        data={**ORDER,'channel':'스마트스토어','amount':'10000','shipping':'0'}
        with self.s.db:
            self.s.db.execute('INSERT INTO smartstore_erp_rows(id,identity,day,data,components,issue) VALUES(?,?,?,?,?,?)',
                              ('UNMAPPED','UNMAPPED',DAY,encode(data),'[]','상품 매칭 필요'))
        root=tk.Tk();root.withdraw()
        try:
            with patch.dict(os.environ,{'REQM_SELF_TEST':'1'}):app=Desktop(root,self.s)
            panel=app.voucher_panel;panel.start.set(DAY);panel.end.set(DAY);panel.refresh()
            key=panel.tree.get_children()[0]
            self.assertEqual(panel.tree.set(key,'품목명'),data['product'])
            self.assertEqual(panel.tree.set(key,'수량'),data['quantity'])
            panel.tree.selection_set(key)
            with patch.object(app,'edit_erp_set') as mapping:
                panel.edit();mapping.assert_called_once()
        finally:root.destroy()
