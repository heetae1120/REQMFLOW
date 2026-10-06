"""Add Ecount item-master rows missing from REQM FLOW's local reference CSV."""
from __future__ import annotations

import argparse
import csv
from pathlib import Path

from openpyxl import load_workbook


FIELDS = (
    "item_code", "representative_name", "alias_count", "first_source_row",
    "review_status", "is_active",
)


def source_items(path: Path) -> list[tuple[int, str, str]]:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = book.active
        sheet.reset_dimensions()
        rows = list(sheet.iter_rows(values_only=True))
    finally:
        book.close()
    if not rows or len(rows[0]) < 2:
        raise ValueError("품목 엑셀의 품목코드·품목명 열을 찾을 수 없습니다.")
    result = []
    seen = set()
    for excel_row, row in enumerate(rows[1:], start=2):
        code = str(row[0] or "").strip()
        name = str(row[1] or "").strip()
        if not code:
            continue
        key = code.casefold()
        if key in seen:
            raise ValueError(f"품목 엑셀에 중복 코드가 있습니다: {code}")
        if not name:
            raise ValueError(f"품목명이 비어 있습니다: 엑셀 {excel_row}행 {code}")
        seen.add(key)
        result.append((excel_row, code, name))
    return result


def sync(source: Path, target: Path) -> list[dict[str, str]]:
    with target.open("r", encoding="utf-8-sig", newline="") as handle:
        existing = list(csv.DictReader(handle))
    existing_codes = {row["item_code"].strip().casefold() for row in existing}
    added = []
    for excel_row, code, name in source_items(source):
        key = code.casefold()
        if key in existing_codes:
            continue
        existing_codes.add(key)
        added.append({
            "item_code": code,
            "representative_name": name,
            "alias_count": "0",
            "first_source_row": str(excel_row),
            "review_status": "confirmed",
            "is_active": "true",
        })
    temporary = target.with_suffix(target.suffix + ".tmp")
    with temporary.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=FIELDS, lineterminator="\n")
        writer.writeheader()
        writer.writerows(existing + added)
    temporary.replace(target)
    return added


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument(
        "--target", type=Path,
        default=Path(__file__).resolve().parents[1] / "supabase/ecount_migration/data/ecount_item_reference.csv",
    )
    args = parser.parse_args()
    added = sync(args.source, args.target)
    print(f"추가 {len(added):,}건")
    print(f"총 {sum(1 for _ in csv.DictReader(args.target.open(encoding='utf-8-sig'))):,}건")


if __name__ == "__main__":
    main()
