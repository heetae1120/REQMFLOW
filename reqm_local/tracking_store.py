"""Persist daily source files, remote evidence, reviews and invoice confirmations."""
from __future__ import annotations

import hashlib
import getpass
import json
import uuid
from datetime import date, datetime
from pathlib import Path

from .files import parse_orders, workbook_bytes
from .tracking import HEADERS, STATES, delivery_key, identity, manual_invoice, reconcile


def encode(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def valid_day(value):
    return date.fromisoformat(str(value)).isoformat()


def now():
    return datetime.now().isoformat(timespec='seconds')


class TrackingOperations:
    def initialize_tracking(self):
        self.db.executescript('''
        CREATE TABLE IF NOT EXISTS tracking_jobs(
            id TEXT PRIMARY KEY, identity TEXT UNIQUE NOT NULL, day TEXT NOT NULL,
            kind TEXT NOT NULL, source_name TEXT NOT NULL, source_content BLOB NOT NULL,
            created_by TEXT NOT NULL, created_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS tracking_rows(
            id TEXT PRIMARY KEY, job_id TEXT NOT NULL REFERENCES tracking_jobs(id),
            order_id TEXT REFERENCES orders(id) ON DELETE SET NULL, data TEXT NOT NULL,
            state TEXT NOT NULL, tracking TEXT NOT NULL DEFAULT '', reason TEXT NOT NULL DEFAULT '',
            evidence TEXT NOT NULL DEFAULT '[]', queried_at TEXT NOT NULL DEFAULT '',
            confirmed_at TEXT NOT NULL DEFAULT '', confirmed_by TEXT NOT NULL DEFAULT '');
        CREATE TABLE IF NOT EXISTS tracking_queries(
            id TEXT PRIMARY KEY, job_id TEXT NOT NULL REFERENCES tracking_jobs(id),
            from_day TEXT NOT NULL, through_day TEXT NOT NULL, remote_rows TEXT NOT NULL,
            queried_at TEXT NOT NULL);
        ''')

    def tracking_actor(self):
        return self.settings.get('cloud_email') or getpass.getuser()

    def import_tracking_files(self, paths, on, kind='b2c', channel_override=None):
        on = valid_day(on)
        if kind not in ('b2c', 'b2c_buying', 'b2b'):
            raise ValueError('위킵 주문 유형을 선택하세요.')
        profiles = [p for p in self.settings['profiles'] if p.get('purpose', 'order') == 'order']
        prepared = []
        for path in paths:
            source = Path(path)
            content = source.read_bytes()
            rows = parse_orders(source, profiles, channel_override=channel_override)
            if not rows:
                raise ValueError(f'{source.name}: 입력할 주문이 없습니다.')
            # Include parsed mapping so a different marketplace cannot silently reuse a job.
            key = hashlib.sha256(encode([on, kind, hashlib.sha256(content).hexdigest(), rows]).encode()).hexdigest()
            prepared.append((source.name, content, rows, key))
        jobs, duplicates = [], 0
        existing_orders = [(r['id'], json.loads(r['data'])) for r in self.db.execute('SELECT id,data FROM orders')]
        with self.db:
            for name, content, rows, key in prepared:
                previous = self.db.execute('SELECT id FROM tracking_jobs WHERE identity=?', (key,)).fetchone()
                if previous:
                    duplicates += 1
                    jobs.append(previous['id'])
                    continue
                job = uuid.uuid4().hex
                self.db.execute('INSERT INTO tracking_jobs VALUES(?,?,?,?,?,?,?,?)',
                                (job, key, on, kind, name, content, self.tracking_actor(), now()))
                for data in rows:
                    matches = [oid for oid, old in existing_orders if
                               all(identity(data.get(k)) == identity(old.get(k)) for k in ('channel', 'account', 'order_no'))
                               and identity(data.get('order_no'))
                               and (identity(data.get('line_no')) == identity(old.get('line_no')) if data.get('line_no') else
                                    all(identity(data.get(k)) == identity(old.get(k)) for k in ('product', 'option', 'quantity')))
                               and delivery_key(data) is not None and delivery_key(data) == delivery_key(old)]
                    self.db.execute('INSERT INTO tracking_rows(id,job_id,order_id,data,state) VALUES(?,?,?,?,?)',
                                    (uuid.uuid4().hex, job, matches[0] if len(matches) == 1 else None,
                                     encode(data), 'pending'))
                jobs.append(job)
            self.event('송장 작업 파일 입력', f'기준일 {on}, 작업 {len(jobs)}, 중복 파일 {duplicates}')
        return jobs, duplicates

    def tracking_jobs(self, on=None):
        rows = self.db.execute('SELECT * FROM tracking_jobs ' + ('WHERE day=? ' if on else '') +
                               'ORDER BY created_at DESC,rowid DESC', (on,) if on else ())
        result = []
        for row in rows:
            job = dict(row)
            job.pop('source_content')
            entries = self.tracking_rows(job['id'])
            job['total'] = len(entries)
            job['confirmed'] = sum(r['state'] == 'confirmed' for r in entries)
            job['channels'] = ', '.join(dict.fromkeys(r['data'].get('channel', '') for r in entries))
            result.append(job)
        return result

    def tracking_rows(self, job_id):
        result = []
        for row in self.db.execute('SELECT * FROM tracking_rows WHERE job_id=? ORDER BY rowid', (job_id,)):
            entry = dict(row)
            entry['data'] = json.loads(entry['data'])
            entry['evidence'] = json.loads(entry['evidence'])
            result.append(entry)
        return result

    def apply_tracking_query(self, job_id, remote_rows, from_day, through_day, expected=None):
        start, end = valid_day(from_day), valid_day(through_day)
        if start > end:
            raise ValueError('위킵 등록 시작일이 종료일보다 늦습니다.')
        entries = self.tracking_rows(job_id)
        if not entries:
            raise ValueError('조회할 출고 작업을 찾지 못했습니다.')
        # A background browser operation must not replace changes made in another PC.
        if expected is not None and encode(entries) != expected:
            raise ValueError('조회 중 작업이 변경되었습니다. 최신 작업에서 다시 조회하세요.')
        results = reconcile([r['data'] for r in entries], remote_rows)
        stamp = now()
        with self.db:
            self.db.execute('INSERT INTO tracking_queries VALUES(?,?,?,?,?,?)',
                            (uuid.uuid4().hex, job_id, start, end, encode(remote_rows), stamp))
            for entry, result in zip(entries, results):
                if entry['state'] == 'confirmed':
                    reason = entry['reason']
                    if result['state'] != 'matched' or result['tracking'] != entry['tracking']:
                        reason = '확정 송장 유지 · 재조회 결과 확인 필요: ' + result['reason']
                    self.db.execute('UPDATE tracking_rows SET evidence=?,queried_at=?,reason=? WHERE id=?',
                                    (encode(result['candidates']), stamp, reason, entry['id']))
                    continue
                self.db.execute('UPDATE tracking_rows SET state=?,tracking=?,reason=?,evidence=?,queried_at=? WHERE id=?',
                                (result['state'], result['tracking'], result['reason'],
                                 encode(result['candidates']), stamp, entry['id']))
            self.event('위킵 송장 조회', f'작업 {job_id}, 등록일 {start}~{end}, 원격 {len(remote_rows)}행')

    def confirm_tracking(self, job_id, row_ids, manual=None, reason=''):
        entries = {r['id']: r for r in self.tracking_rows(job_id)}
        row_ids = list(dict.fromkeys(row_ids))
        if not row_ids or any(rid not in entries for rid in row_ids):
            raise ValueError('확정할 주문을 선택하세요.')
        value = manual_invoice(manual) if manual is not None else None
        if value and not reason.strip():
            raise ValueError('수동 연결 사유를 입력하세요.')
        for rid in row_ids:
            row = entries[rid]
            if value is None and row['state'] not in ('matched', 'confirmed'):
                raise ValueError('자동 매칭 완료 주문만 확정할 수 있습니다. 검토 주문은 수동 연결하세요.')
            if row['state'] == 'confirmed' and value and value != row['tracking']:
                raise ValueError('이미 확정된 송장은 변경할 수 없습니다. 기존 출고정보를 확인하세요.')
        with self.db:
            for rid in row_ids:
                row = entries[rid]
                tracking = value or row['tracking']
                if row['order_id']:
                    order = self.db.execute('SELECT data FROM orders WHERE id=?', (row['order_id'],)).fetchone()
                    if order:
                        data = json.loads(order['data'])
                        if data.get('tracking') and data['tracking'] != tracking:
                            raise ValueError('연결된 FLOW 주문에 다른 송장이 있습니다. 기존 출고정보를 확인하세요.')
                        data.update(tracking=tracking, _wekeep_tracking=tracking)
                        self.db.execute('UPDATE orders SET data=? WHERE id=?', (encode(data), row['order_id']))
                detail = ('작업자 수동 확인: ' + reason.strip()) if value else row['reason']
                self.db.execute('UPDATE tracking_rows SET state=?,tracking=?,reason=?,confirmed_at=?,confirmed_by=? WHERE id=?',
                                ('confirmed', tracking, detail, now(), self.tracking_actor(), rid))
            self.event('송장 확정', f'작업 {job_id}, {len(row_ids)}행, 사유 {reason}')

    def export_tracking(self, job_id, path, confirmed_only=False):
        job_ids = [job_id] if isinstance(job_id, str) else list(dict.fromkeys(job_id))
        entries = [r for jid in job_ids for r in self.tracking_rows(jid)]
        if not entries:
            raise ValueError('내려받을 출고 작업을 선택하세요.')
        rows = [r for r in entries if not confirmed_only or r['state'] == 'confirmed']
        if not rows:
            raise ValueError('확정된 송장 주문이 없습니다.')
        # Preserve the Shipping program's eleven-column layout, invoice at J.
        values = []
        for index, row in enumerate(rows, 1):
            d = row['data']
            values.append([identity(d.get(k)) for k in ('order_no', 'channel', 'product', 'quantity',
                           'recipient', 'phone', 'postcode', 'address', 'memo')] +
                          [row['tracking'] if row['state'] == 'confirmed' else '', str(index)])
        content = workbook_bytes(HEADERS, values, '택배출고')
        # All unresolved rows remain visible without an unconfirmed invoice.
        from io import BytesIO
        from openpyxl import load_workbook
        book = load_workbook(BytesIO(content))
        status = book.create_sheet('송장 확인 상태')
        status.append(['주문번호', '판매처', '수령인', '상태', '확인 내용', '송장번호', '확정 작업자', '확정 시각'])
        for row in rows:
            d = row['data']
            status.append([identity(d.get('order_no')), identity(d.get('channel')), identity(d.get('recipient')),
                           STATES[row['state']], row['reason'], row['tracking'] if row['state'] == 'confirmed' else '',
                           row['confirmed_by'], row['confirmed_at']])
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
        for sheet in book:
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
            sheet.row_dimensions[1].height = 26
            for cell in sheet[1]:
                cell.font = Font(color='FFFFFF', bold=True)
                cell.fill = PatternFill('solid', fgColor='173C4A')
            for cells in sheet.iter_rows(min_row=2):
                for cell in cells:
                    if isinstance(cell.value, str):
                        cell.data_type = 's'
                        cell.number_format = '@'
                    cell.alignment = Alignment(vertical='top', wrap_text=True)
            for i, cell in enumerate(sheet[1], 1):
                sheet.column_dimensions[get_column_letter(i)].width = 60 if cell.value in ('주소', '확인 내용') else 24
        output = BytesIO()
        book.save(output)
        book.close()
        Path(path).write_bytes(output.getvalue())
        with self.db:
            self.event('송장 엑셀 다운로드', f'작업 {job_id}, {len(rows)}행, 확정건만 {confirmed_only}')

    def export_tracking_day(self, on, path, confirmed_only=False):
        jobs = self.tracking_jobs(valid_day(on))
        self.export_tracking([j['id'] for j in reversed(jobs)], path, confirmed_only)

    def export_tracking_source(self, job_id, path):
        job = self.db.execute('SELECT * FROM tracking_jobs WHERE id=?', (job_id,)).fetchone()
        if not job:
            raise ValueError('출고 작업을 찾지 못했습니다.')
        if Path(job['source_name']).suffix.lower() not in ('.xlsx', '.xlsm', '.csv'):
            raise ValueError('원본 송장 반영은 XLSX, XLSM, CSV를 지원합니다. 공통 송장 엑셀을 사용하세요.')
        entries = [{'data': r['data'], 'tracking': r['tracking']} for r in self.tracking_rows(job_id)
                   if r['state'] == 'confirmed']
        content = self._update_source_workbook(job['source_name'], job['source_content'], entries) if entries else job['source_content']
        Path(path).write_bytes(content)
