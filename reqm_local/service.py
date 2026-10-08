from __future__ import annotations

import hashlib
import io
import json
import os
import sqlite3
import tempfile
import time
import uuid
import re
import csv
from collections import Counter
from difflib import SequenceMatcher
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from openpyxl import load_workbook

from ecount_sales_core import ReferenceCatalog, SmartStoreOrder, convert_orders, normalize_source
from .files import (
    parse_orders, read_rows, identifier, workbook_bytes, wekeep_workbook_bytes,
    REQUEST_COLUMNS, WEKEEP_REQUEST_COLUMNS, settings_at,
)
from .esm import parse_esm_rows
from .profiles import SMARTSTORE_PURCHASE_CHANNEL


SMARTSTORE_CHANNEL = '리큐엠_스마트스토어'
SMARTSTORE_ERP_CHANNEL = '리큐엠_스마트스토어_ERP'


def is_smartstore_channel(value):
    return identifier(value).replace(' ', '') in (
        '스마트스토어', SMARTSTORE_CHANNEL, SMARTSTORE_ERP_CHANNEL, SMARTSTORE_PURCHASE_CHANNEL,
    )


def display_channel(value):
    text=identifier(value).replace(' ', '')
    if is_smartstore_channel(text):
        return '스마트스토어'
    return identifier(value).removeprefix('리큐엠_')


def is_excluded_channel(value, excluded_channels=None):
    excluded={identifier(item).replace(' ', '').removeprefix('리큐엠_') for item in (excluded_channels or [])}
    return display_channel(value).replace(' ', '') in excluded


def number(value, positive=False):
    try:
        result = Decimal(str(value).replace(',', ''))
    except InvalidOperation:
        raise ValueError(f'숫자가 아닌 값: {value!r}') from None
    if not result.is_finite() or result < 0 or (positive and result <= 0):
        raise ValueError(f'허용되지 않는 금액/수량: {value!r}')
    return result


def money(value):
    try:return number(value or '0')
    except ValueError:
        matches=re.findall(r'\d[\d,]*(?:\.\d+)?',identifier(value))
        if len(matches)==1:return number(matches[0])
        raise


def quantity(value):
    result = number(value, True)
    if result != result.to_integral_value():
        raise ValueError('출고 수량은 양의 정수여야 합니다.')
    return int(result)


def day(value):
    value = identifier(value)[:10].replace('.', '-').replace('/', '-')
    if len(value) == 8 and value.isdigit():
        value = f'{value[:4]}-{value[4:6]}-{value[6:]}'
    return date.fromisoformat(value).isoformat()


def encode(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True)


def warehouse_code(value):
    text = identifier(value).replace(' ', '')
    if text.startswith('100'):
        return '100'
    if text.startswith('300'):
        return '300'
    raise ValueError('출하창고는 100 본사창고 또는 300 위킵창고만 사용할 수 있습니다.')


class Operations:
    """All mutations are transactional. No GUI, network, or user session dependencies."""

    def __init__(self, folder, reference_dir):
        self.folder = Path(folder)
        self.reference_dir = Path(reference_dir)
        self.folder.mkdir(parents=True, exist_ok=True)
        self.settings = settings_at(self.folder)
        self.catalog = ReferenceCatalog.from_csv_dir(self.reference_dir)
        self.db = sqlite3.connect(self.folder / 'operations.sqlite3')
        self.db.row_factory = sqlite3.Row
        self.db.execute('PRAGMA foreign_keys=ON')
        self.db.execute('PRAGMA journal_mode=WAL')
        self.db.executescript('''
        CREATE TABLE IF NOT EXISTS orders(id TEXT PRIMARY KEY, identity TEXT UNIQUE NOT NULL, data TEXT NOT NULL,
            components TEXT NOT NULL DEFAULT '[]', state TEXT NOT NULL, issue TEXT NOT NULL DEFAULT '');
        CREATE TABLE IF NOT EXISTS mappings(key TEXT PRIMARY KEY, components TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS item_names(erp_item_code TEXT PRIMARY KEY, matching_name TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS event_rules(id TEXT PRIMARY KEY, channel TEXT NOT NULL, name TEXT NOT NULL,
            source_product TEXT NOT NULL, source_option TEXT NOT NULL, target_product TEXT NOT NULL,
            target_option TEXT NOT NULL, target_amount TEXT NOT NULL, components TEXT NOT NULL DEFAULT '[]',
            active INTEGER NOT NULL DEFAULT 1,
            UNIQUE(channel, source_product, source_option));
        CREATE TABLE IF NOT EXISTS requests(id TEXT PRIMARY KEY, day TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS request_lines(id TEXT PRIMARY KEY, request_id TEXT REFERENCES requests(id),
            order_id TEXT REFERENCES orders(id), component INTEGER, qty INTEGER NOT NULL,
            UNIQUE(order_id, component));
        CREATE TABLE IF NOT EXISTS shipments(id TEXT PRIMARY KEY, line_id TEXT REFERENCES request_lines(id),
            qty INTEGER NOT NULL, tracking TEXT NOT NULL, day TEXT NOT NULL, erp_id TEXT, erp_amount TEXT,
            UNIQUE(line_id,tracking,day));
        CREATE TABLE IF NOT EXISTS smartstore_erp_rows(id TEXT PRIMARY KEY, identity TEXT UNIQUE NOT NULL,
            day TEXT NOT NULL, data TEXT NOT NULL, components TEXT NOT NULL DEFAULT '[]',
            erp_id TEXT, issue TEXT NOT NULL DEFAULT '');
        CREATE TABLE IF NOT EXISTS smartstore_purchase_rows(id TEXT PRIMARY KEY, identity TEXT UNIQUE NOT NULL,
            day TEXT NOT NULL, data TEXT NOT NULL, smartstore_erp_id TEXT);
        CREATE TABLE IF NOT EXISTS esm_erp_rows(id TEXT PRIMARY KEY, identity TEXT UNIQUE NOT NULL,
            day TEXT NOT NULL, data TEXT NOT NULL, components TEXT NOT NULL DEFAULT '[]',
            erp_id TEXT, issue TEXT NOT NULL DEFAULT '');
        CREATE TABLE IF NOT EXISTS fees(bundle TEXT PRIMARY KEY, erp_id TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS artifacts(id TEXT PRIMARY KEY, kind TEXT NOT NULL, day TEXT NOT NULL,
            content BLOB NOT NULL, registered INTEGER NOT NULL DEFAULT 0);
        CREATE TABLE IF NOT EXISTS source_files(name TEXT PRIMARY KEY, extension TEXT NOT NULL,
            content BLOB NOT NULL, updated_at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS artifact_notes(artifact_id TEXT NOT NULL, row_key TEXT NOT NULL,
            note TEXT NOT NULL DEFAULT '', PRIMARY KEY(artifact_id,row_key));
        CREATE TABLE IF NOT EXISTS events(id INTEGER PRIMARY KEY, at TEXT NOT NULL, action TEXT NOT NULL, detail TEXT NOT NULL);
        ''')
        event_columns = {row['name'] for row in self.db.execute('PRAGMA table_info(event_rules)')}
        if 'components' not in event_columns:
            self.db.execute("ALTER TABLE event_rules ADD COLUMN components TEXT NOT NULL DEFAULT '[]'")
        shipment_columns = {row['name'] for row in self.db.execute('PRAGMA table_info(shipments)')}
        if 'erp_amount' not in shipment_columns:
            self.db.execute("ALTER TABLE shipments ADD COLUMN erp_amount TEXT")
        for name, definition in [('erp_code', 'TEXT'), ('erp_name', 'TEXT'), ('erp_warehouse', 'TEXT'), ('amount_split_confirmed', 'INTEGER NOT NULL DEFAULT 0')]:
            if name not in shipment_columns:
                self.db.execute(f'ALTER TABLE shipments ADD COLUMN {name} {definition}')

    def close(self):
        self.db.close()

    def event(self, action, detail):
        self.db.execute('INSERT INTO events(at,action,detail) VALUES(?,?,?)', (datetime.now().isoformat(timespec='seconds'), action, detail))

    def _write_settings(self):
        path = self.folder / 'settings.json'
        temporary = path.with_suffix('.json.tmp')
        temporary.write_text(json.dumps(self.settings,ensure_ascii=False,indent=2),encoding='utf-8')
        for attempt in range(4):
            try:
                os.replace(temporary,path)
                return
            except PermissionError:
                if attempt == 3:
                    raise
                time.sleep(.03 * (attempt + 1))

    def channel_customer_code(self, channel):
        from .shipping import channel_key
        key=channel_key(channel)
        saved={identifier(code) for name,code in self.settings.get('channel_customer_codes',{}).items() if channel_key(name)==key and identifier(code)}
        if saved:return next(iter(saved)) if len(saved)==1 else ''
        codes={identifier(row.get('ecount_customer_code')) for name,row in self.catalog.channels.items()
               if key in (channel_key(name),channel_key(row.get('normalized_name'))) and identifier(row.get('ecount_customer_code'))}
        return next(iter(codes)) if len(codes)==1 else ''

    def catalog_channel(self, channel):
        from .shipping import channel_key
        key=channel_key(channel)
        code=self.channel_customer_code(channel)
        return next((name for name,row in self.catalog.channels.items() if channel_key(name)==key
                     and identifier(row.get('ecount_customer_code'))==code),channel)

    def conversion_catalog(self, channel):
        """Reuse equivalent marketplace keys without changing the shared reference DB."""
        from copy import copy
        from .shipping import channel_key
        catalog = copy(self.catalog)
        key = channel_key(channel)
        for attribute in ('mappings', 'price_templates'):
            original = getattr(self.catalog, attribute)
            merged = dict(original)
            candidates = {}
            for (source_channel, source), value in original.items():
                if channel_key(source_channel) == key:
                    candidates.setdefault(source, []).append(value)
            for source, values in candidates.items():
                target = (channel, source)
                if target not in merged and all(value == values[0] for value in values):
                    merged[target] = values[0]
            setattr(catalog, attribute, merged)
        return catalog

    def customer_choices(self):
        choices = {}
        for channel, values in self.catalog.channels.items():
            code = identifier(values.get('ecount_customer_code',''))
            if code:
                name = identifier(values.get('ecount_customer_name',''))
                choices[code] = f'{code} · {name or channel}'
        for channel, code in self.settings.get('channel_customer_codes',{}).items():
            code = identifier(code)
            if code and code not in choices:
                choices[code] = f'{code} · {channel}'
        return [choices[code] for code in sorted(choices)]

    def _prepare_component_definitions(self, channel, components):
        sources = [dict(component) for component in components]
        customer = next((identifier(item.get('customer','')) for item in sources if identifier(item.get('customer',''))), '')
        customer = customer or self.channel_customer_code(channel)
        if customer:
            for item in sources:
                if not identifier(item.get('customer','')):
                    item['customer'] = customer
        validated = self._validate_component_definitions(sources, channel)
        return validated, customer

    def mapping_key(self, data):
        return encode([data['channel'], data['account'], data['product'], data['option']])

    def bundle_key(self, data):
        return encode([data['channel'], data['account'], data['bundle'] or data['order_no']])

    def duplicate_key(self, data):
        return tuple(identifier(data.get(key, '')) for key in ('channel','order_no','line_no','product','option','quantity','amount'))

    def _event_source(self, data):
        return (
            data.get('source_product', data.get('product', '')),
            data.get('source_option', data.get('option', '')),
            data.get('source_amount', data.get('amount', '')),
        )

    def _apply_event_rule(self, data):
        source_product, source_option, source_amount = self._event_source(data)
        rule = self.db.execute('''SELECT * FROM event_rules WHERE active=1 AND channel=?
            AND source_product=? AND source_option=?''', (data['channel'],source_product,source_option)).fetchone()
        if not rule:
            return data
        data['source_product'], data['source_option'], data['source_amount'] = source_product, source_option, source_amount
        data['product'] = rule['target_product'] or source_product
        data['option'] = rule['target_option'] or source_option
        data['amount'] = str(number(rule['target_amount']) * quantity(data['quantity'])) if rule['target_amount'] else source_amount
        data['event_name'], data['event_rule_id'] = rule['name'], rule['id']
        data['event_unit_amount'] = rule['target_amount']
        data['event_components'] = json.loads(rule['components'] or '[]')
        return data

    def _validate_component_definitions(self, components, channel=''):
        if not components:
            raise ValueError('구성품을 한 개 이상 입력하세요.')
        validated = []
        for source in components:
            component = {key:identifier(source.get(key,'')) for key in ('code','logistics_code','name','warehouse','customer')}
            labels = {'code':'ERP 품목코드','logistics_code':'물류사 품목코드','name':'품목명','warehouse':'창고','customer':'ERP 거래처코드'}
            for key in ('code','logistics_code','name','warehouse','customer'):
                if not component[key]:
                    if key == 'customer' and channel:
                        raise ValueError(f'{channel}의 ERP 거래처코드를 입력하세요. 한 번 저장하면 같은 판매처의 다음 매칭에 자동 적용됩니다.')
                    raise ValueError(f'구성품의 {labels[key]}를 입력하세요.')
            component['warehouse'] = warehouse_code(component['warehouse'])
            component['quantity'] = quantity(source.get('quantity',''))
            component['unit_amount'] = str(number(source.get('unit_amount','0')))
            if number(component['unit_amount']) != number(component['unit_amount']).to_integral_value():
                raise ValueError('구성품 단가는 원 단위 정수로 입력하세요.')
            validated.append(component)
        return validated

    def _build_components(self, data, components):
        total = number(data['amount'])
        source_quantity = quantity(data['quantity'])
        # FLOW receives the final quantity from each marketplace spreadsheet.
        # Component quantities are descriptive legacy data only and must never
        # multiply the requested shipping quantity.
        fixed = sum(number(c['unit_amount']) * source_quantity for c in components[1:])
        if fixed > total:
            return [], '부속품 금액이 주문 금액을 초과합니다.'
        result = []
        for index, comp in enumerate(components):
            q = source_quantity
            amount = total - fixed if index == 0 else number(comp['unit_amount']) * q
            result.append({**comp, 'quantity': q, 'source_quantity': source_quantity, 'amount': str(amount)})
        return self._attach_wekeep_skus(data, result)

    def _normalize_delivery(self, data):
        original={key:identifier(data.get(key,'')) for key in ('recipient','phone','postcode','address','memo')}
        data.update(original)
        digits=re.sub(r'\D','',data['phone'])
        if digits.startswith('82') and digits[2:].startswith('0'):
            digits=digits[2:]
        if 9 <= len(digits) <= 12:
            data['phone'] = digits
        elif digits:
            data['phone'] = ''
        postcode=re.sub(r'\D','',data['postcode'])
        data['postcode']=postcode[:5] if len(postcode)>=5 else ''
        data['address']=' '.join(data['address'].split())
        changes={key:{'original':original[key],'normalized':data[key]} for key in original if original[key]!=data[key]}
        data['delivery_changes']=changes
        return data

    def item_candidates(self, data, limit=8):
        from .shipping import compact, channel_key, _component_rows
        source=compact(f"{data.get('product','')} {data.get('option','')}")
        shipping=self._shipping_catalog()
        alias_codes=set()
        if shipping:
            aliases=shipping.aliases.get((channel_key(data.get('channel')),source),[]) or shipping.aliases.get(('',source),[])
            alias_codes={identifier(component.get('item_code')) for alias in aliases for component in _component_rows(alias.get('components'))}
        combined={}
        for code,item in self.catalog.items.items():
            combined[str(code)]={'code':str(code),'name':identifier(item.get('representative_name') or item.get('item_name'))}
        if shipping:
            for key,item in shipping.items.items():
                code=identifier(item.get('item_code')) or key
                combined.setdefault(code,{'code':code,'name':identifier(item.get('standard_name'))})
        results=[]
        for row in combined.values():
            target=compact(f"{row['code']} {row['name']}")
            score=SequenceMatcher(None,source,target).ratio() if source and target else 0
            source_tokens=set(re.findall(r'[0-9a-z]+|[가-힣]+',source))
            target_tokens=set(re.findall(r'[0-9a-z]+|[가-힣]+',target))
            overlap=len(source_tokens & target_tokens)/max(1,len(source_tokens | target_tokens))
            score=max(score,overlap)
            if score < .28:continue
            sku=bool(shipping and shipping.sku_for(row['code']))
            reason='상품명·옵션 유사' if score<.85 else '상품명·옵션 높은 일치'
            results.append({**row,'score':round(score,3),'reason':reason,'sku':sku,'alias':row['code'] in alias_codes})
        return sorted(results,key=lambda row:(row['score'],row['sku']),reverse=True)[:limit]

    def _shipping_catalog(self):
        return getattr(self.catalog, 'shipping_catalog', None)

    def _attach_wekeep_skus(self, data, components):
        """Require a verified WeKeep SKU whenever the shared shipping DB is available."""
        shipping = self._shipping_catalog()
        if shipping is None:
            return components, ''
        if not shipping.available:
            return [], f'{shipping.blocking_reason} 안전을 위해 출고를 중단합니다.'
        result = []
        source_quantity = quantity(data['quantity'])
        for component in components:
            item_code = identifier(component.get('code') or component.get('logistics_code'))
            mapping = shipping.sku_for(item_code)
            if not mapping:
                return [], f'위킵 SKU 미등록: {item_code or "품목코드 없음"}. 출고를 중단합니다.'
            sku_no = identifier(mapping.get('sku_no'))
            product_name = identifier(mapping.get('product_name'))
            if not sku_no or not product_name:
                return [], f'위킵 SKU 정보 누락: {item_code}. 출고를 중단합니다.'
            result.append({
                **component,
                'quantity': source_quantity,
                'source_quantity': source_quantity,
                'name': product_name,
                'wekeep_manage_code': identifier(mapping.get('wekeep_manage_code')) or item_code,
                'wekeep_product_name': product_name,
                'sku_no': sku_no,
                'customer_barcode': identifier(mapping.get('customer_barcode')),
            })
        return result, ''

    def _match_shipping_catalog(self, data):
        shipping = self._shipping_catalog()
        if shipping is None:
            return None
        if not shipping.available:
            return [], f'{shipping.blocking_reason} 안전을 위해 출고를 중단합니다.'
        resolved = shipping.resolve(data)
        if resolved['state'] == 'no_match':
            return None
        if resolved['state'] != 'ready':
            return [], resolved['reason']
        customer = self.channel_customer_code(data['channel'])
        if not customer:
            return [], f'{data["channel"]}의 ERP 거래처코드가 없어 출고를 중단합니다.'
        definitions = [
            {
                'code': component['item_code'],
                'logistics_code': component['item_code'],
                'name': component['wekeep_product_name'],
                'quantity': 1,
                'unit_amount': '0',
                'warehouse': self.settings['warehouse'],
                'customer': customer,
                'match_method': resolved['method'],
            }
            for component in resolved['components']
        ]
        return self._build_components(data, definitions)

    def match(self, data):
        if any(word in data['status'] for word in ('취소', '반품', '교환')):
            return [], '취소·반품·교환 주문: 출고 보류'
        if not all(data[k] for k in ('recipient', 'phone', 'address', 'postcode')):
            return [], '배송정보 누락: 수령인·연락처·우편번호·주소 확인'
        if data.get('event_components'):
            return self._build_components(data, [{**c,'match_method':'이벤트 변환 규칙'} for c in data['event_components']])
        custom = self.db.execute('SELECT components FROM mappings WHERE key=?', (self.mapping_key(data),)).fetchone()
        if custom:
            return self._build_components(data, [{**c,'match_method':'플로우 저장 매칭'} for c in json.loads(custom[0])])
        shipping_match = self._match_shipping_catalog(data)
        if shipping_match is not None:
            return shipping_match
        catalog_channel=self.catalog_channel(data['channel'])
        if catalog_channel not in self.catalog.channels:
            return [], '판매처 ERP 거래처 연결 필요: 상품 매칭에서 거래처코드를 입력하세요.'
        order = SmartStoreOrder(0, data['order_no'], data['line_no'], datetime.now(), data['status'],
            data['product'], data['option'], Decimal(data['quantity']), number(data['amount']), include_shipping=False)
        result = convert_orders([order], self.conversion_catalog(catalog_channel), channel_name=catalog_channel, default_warehouse=self.settings['warehouse'])
        if result.issues or not result.is_reconciled:
            return [], '; '.join(x.reason for x in result.issues) or '금액 배분 불일치'
        grouped = {}
        for line in result.lines:
            if not line.item_code or line.item_code not in self.catalog.items:
                return [], '등록된 품목코드가 필요합니다.'
            key = (line.item_code, line.warehouse)
            if key not in grouped:
                grouped[key] = dict(code=line.item_code, logistics_code=line.item_code, name=line.item_name,
                    quantity=0, amount='0', warehouse=line.warehouse, customer=line.customer_code)
            component = grouped[key]
            component['quantity'] += quantity(line.quantity)
            component['amount'] = str(number(component['amount']) + line.total)
            component['match_method'] = '판매전표 DB 변환 규칙'
        return self._attach_wekeep_skus(data, list(grouped.values()))

    def import_files(self, paths, channel_override=None, registered_on=None):
        incoming = []
        registered_on=day(registered_on or date.today().isoformat())
        profiles = [profile for profile in self.settings['profiles'] if profile.get('purpose','order') == 'order']
        for path in paths:
            incoming.extend(parse_orders(path, profiles, channel_override=channel_override))
        inserted = duplicates = 0
        with self.db:
            for path in paths:
                source=Path(path)
                self.db.execute(
                    'INSERT OR REPLACE INTO source_files(name,extension,content,updated_at) VALUES(?,?,?,?)',
                    (source.name,source.suffix.lower(),source.read_bytes(),datetime.now().isoformat(timespec='seconds')),
                )
            for data in incoming:
                data['_registered_day']=registered_on
                seed = f"{data['source_file']}|{data['source_row']}|{data.get('product','')}"
                token = hashlib.sha256(seed.encode('utf-8')).hexdigest()[:12].upper()
                data['order_no'] = data.get('order_no') or f'AUTO-ORDER-{token}'
                data['line_no'] = data.get('line_no') or f'AUTO-LINE-{token}'
                data['quantity'] = data.get('quantity') or '1'
                data['amount'] = data.get('amount') or '0'
                for field in ('source_item_code','product','option','recipient','phone','postcode','address','memo','paid_at','status','bundle','shipping'):
                    data[field] = data.get(field, '')
                data = self._normalize_delivery(data)
                data = self._apply_event_rule(data)
                data['quantity'] = str(quantity(data['quantity']))
                data['amount'] = str(number(data['amount']))
                data['shipping'] = str(money(data['shipping']))
                if any(number(data[k]) != number(data[k]).to_integral_value() for k in ('amount','shipping')):
                    raise ValueError('현재 원화 금액은 원 단위 정수만 지원합니다.')
                key = encode([data['channel'], data['account'], data['line_no']])
                old = self.db.execute('SELECT data FROM orders WHERE identity=?', (key,)).fetchone()
                if old:
                    if self.settings.get('allow_duplicate_orders', True):
                        key = encode([data['channel'], data['account'], data['line_no'], uuid.uuid4().hex])
                        old = None
                if old:
                    before = json.loads(old[0])
                    compare = lambda x: {k:v for k,v in x.items() if k not in ('source_file','source_row')}
                    if compare(before) != compare(data):
                        raise ValueError(f"상품주문번호 {data['line_no']}: 기존 주문과 내용이 다릅니다. 원본·출고 상태 확인 필요. 이번 입력은 취소됩니다.")
                    duplicates += 1
                    continue
                components, issue = self.match(data)
                self.db.execute('INSERT INTO orders VALUES(?,?,?,?,?,?)', (uuid.uuid4().hex, key, encode(data), encode(components), '검토 필요' if issue else '출고 준비', issue))
                inserted += 1
            self.event('주문 가져오기', f'신규 {inserted}, 중복 {duplicates}')
        return inserted, duplicates

    @staticmethod
    def _tracking_text(values):
        return ', '.join(dict.fromkeys(identifier(value) for value in values if identifier(value)))

    def _update_source_workbook(self, name, content, orders):
        extension=Path(name).suffix.lower()
        if extension in ('.xlsx','.xlsm'):
            book=load_workbook(io.BytesIO(content),keep_vba=extension=='.xlsm')
            sheet=max(book.worksheets,key=lambda value:value.max_row*value.max_column)
            header_row=min(int(order['data'].get('source_header_row') or 1) for order in orders)
            aliases={'송장번호','운송장번호','택배사송장번호'}
            tracking_column=next((cell.column for cell in sheet[header_row] if identifier(cell.value).replace(' ','') in aliases),None)
            if tracking_column is None:
                tracking_column=sheet.max_column+1
                source_header=sheet.cell(header_row,max(1,tracking_column-1))
                target_header=sheet.cell(header_row,tracking_column,'송장번호')
                if source_header.has_style:
                    from copy import copy
                    target_header._style=copy(source_header._style)
                    target_header.font=copy(source_header.font);target_header.fill=copy(source_header.fill)
                    target_header.border=copy(source_header.border);target_header.alignment=copy(source_header.alignment)
                    target_header.number_format=source_header.number_format
            for order in orders:
                target=sheet.cell(int(order['data']['source_row']),tracking_column)
                target.value=order['tracking'];target.data_type='s'
            output=io.BytesIO();book.save(output);book.close()
            return output.getvalue()
        if extension=='.csv':
            text=None
            for encoding in ('utf-8-sig','cp949','utf-8'):
                try:text=content.decode(encoding);break
                except UnicodeDecodeError:continue
            if text is None:raise ValueError(f'{name}: 문자 인코딩을 확인하세요.')
            rows=list(csv.reader(io.StringIO(text)))
            header_row=min(int(order['data'].get('source_header_row') or 1) for order in orders)-1
            aliases={'송장번호','운송장번호','택배사송장번호'}
            tracking_column=next((index for index,value in enumerate(rows[header_row]) if identifier(value).replace(' ','') in aliases),None)
            if tracking_column is None:
                tracking_column=len(rows[header_row]);rows[header_row].append('송장번호')
            for order in orders:
                row_index=int(order['data']['source_row'])-1
                while len(rows[row_index])<=tracking_column:rows[row_index].append('')
                rows[row_index][tracking_column]=order['tracking']
            output=io.StringIO(newline='');csv.writer(output,lineterminator='\r\n').writerows(rows)
            return output.getvalue().encode('utf-8-sig')
        return content

    def _sync_source_tracking(self, order_ids):
        grouped={}
        for order_id in set(order_ids):
            row=self.db.execute('SELECT data FROM orders WHERE id=?',(order_id,)).fetchone()
            if not row:continue
            data=json.loads(row['data'])
            tracking=self._tracking_text(value[0] for value in self.db.execute('''SELECT s.tracking FROM shipments s
                JOIN request_lines r ON s.line_id=r.id WHERE r.order_id=? ORDER BY s.rowid''',(order_id,)))
            data['tracking']=tracking
            self.db.execute('UPDATE orders SET data=? WHERE id=?',(encode(data),order_id))
            if data.get('source_file') and tracking:
                grouped.setdefault(data['source_file'],[]).append({'data':data,'tracking':tracking})
        for name,orders in grouped.items():
            source=self.db.execute('SELECT content FROM source_files WHERE name=?',(name,)).fetchone()
            if not source:continue
            content=self._update_source_workbook(name,source['content'],orders)
            self.db.execute('UPDATE source_files SET content=?,updated_at=? WHERE name=?',(
                content,datetime.now().isoformat(timespec='seconds'),name,
            ))

    def export_source_file(self, name, path):
        row=self.db.execute('SELECT content FROM source_files WHERE name=?',(name,)).fetchone()
        if not row:raise ValueError('워크스페이스에 저장된 판매처 주문 파일을 찾지 못했습니다.')
        Path(path).write_bytes(row['content'])

    def source_file_available(self, name):
        return bool(self.db.execute('SELECT 1 FROM source_files WHERE name=?',(name,)).fetchone())

    def attach_source_file(self, name, path):
        source=Path(path)
        if source.suffix.lower() not in ('.xlsx','.xlsm','.csv'):
            raise ValueError('송장번호 반영은 XLSX, XLSM, CSV 주문 파일을 지원합니다.')
        order_ids=[]
        for row in self.db.execute('SELECT id,data FROM orders'):
            if json.loads(row['data']).get('source_file')==name:order_ids.append(row['id'])
        if not order_ids:raise ValueError('선택한 출력 이력과 연결된 주문을 찾지 못했습니다.')
        with self.db:
            self.db.execute('INSERT OR REPLACE INTO source_files(name,extension,content,updated_at) VALUES(?,?,?,?)',(
                name,source.suffix.lower(),source.read_bytes(),datetime.now().isoformat(timespec='seconds'),
            ))
            self._sync_source_tracking(order_ids)
            self.event('기존 주문 파일 연결',name)

    def match_statistics(self, orders=None):
        orders=list(orders if orders is not None else self.orders())
        ready=[order for order in orders if order['state']=='출고 준비']
        methods=Counter()
        for order in ready:
            method=next((c.get('match_method') for c in order['components'] if c.get('match_method')),'기존 변환 규칙')
            methods[method]+=1
        return {'total':len(orders),'ready':len(ready),'review':sum(o['state']=='검토 필요' for o in orders),
                'rate':round(len(ready)*100/len(orders)) if orders else 0,'methods':dict(methods)}

    def shipping_duplicate_key(self, data):
        from .shipping import compact
        fields=(data.get('order_no'),data.get('recipient'),data.get('phone'),data.get('postcode'),
                data.get('address'),data.get('product'),data.get('option'),data.get('quantity'))
        normalized='|'.join(compact(value) for value in fields)
        return hashlib.sha256(normalized.encode('utf-8')).hexdigest() if data.get('order_no') else ''

    def check_shipping_history(self, client, order_ids):
        selected=[order for order in self.orders() if order['id'] in set(order_ids)]
        numbers=list({order['data']['order_no'] for order in selected if order['data']['order_no']})
        if not numbers:return 0
        response=client.table('shipment_history').select('duplicate_key').in_('order_number',numbers).execute()
        shipped={str(row.get('duplicate_key','')) for row in (response.data or [])}
        found=0
        with self.db:
            for order in selected:
                if self.shipping_duplicate_key(order['data']) not in shipped:continue
                issue='출고 프로그램 이력에 동일 주문·배송정보·상품이 있습니다.'
                current=order['issue']
                if issue not in current:
                    current='; '.join(filter(None,[current,issue]))
                data=order['data'];data['shipping_history_duplicate']=True
                self.db.execute("UPDATE orders SET data=?,state='검토 필요',issue=? WHERE id=?",(encode(data),current,order['id']))
                found+=1
        return found

    def preflight(self, ids):
        selected=[order for order in self.orders() if order['id'] in set(ids)]
        issues=[]
        for order in selected:
            data=order['data'];prefix=f"{data['order_no']} / {data['product']}"
            forced=bool(data.get('force_shipping_approved'))
            if order['state']!='출고 준비' and not forced:issues.append((order['id'],'상태',prefix+' · '+(order['issue'] or order['state'])))
            if order.get('duplicate_count',1)>1 and not forced:issues.append((order['id'],'중복',prefix+' · 현재 입력에 동일 주문행이 반복됩니다.'))
            if not self.channel_customer_code(data['channel']) and not forced:issues.append((order['id'],'전표',prefix+' · ERP 거래처코드가 없습니다.'))
            if not order['components']:issues.append((order['id'],'품목',prefix+' · 출고 품목이 없습니다.'))
            shipping=self._shipping_catalog()
            for component in order['components']:
                if not component.get('name'):issues.append((order['id'],'품목',prefix+' · 출고 품목명이 없습니다.'))
                if not component.get('code') and not forced:issues.append((order['id'],'전표',prefix+' · ERP 품목코드가 없습니다.'))
                if shipping is not None and shipping.available and not component.get('sku_no') and not forced:
                    issues.append((order['id'],'위킵 SKU',prefix+' · 위킵 SKU가 없습니다.'))
        return issues

    def set_force_shipping_approval(self, ids, approved, reason=''):
        ids=list(dict.fromkeys(ids))
        if not ids:raise ValueError('승인할 주문을 선택하세요.')
        rows=self.db.execute(f"SELECT * FROM orders WHERE id IN ({','.join('?' for _ in ids)})",ids).fetchall()
        if len(rows)!=len(ids):raise ValueError('선택한 주문을 다시 확인하세요.')
        if any(row['state'] not in ('검토 필요','출고 준비') for row in rows):
            raise ValueError('출고 요청 전 주문만 강제 출고를 승인하거나 해제할 수 있습니다.')
        reason=identifier(reason)
        if approved and not reason:raise ValueError('강제 출고 승인 사유를 입력하세요.')
        with self.db:
            for row in rows:
                data=json.loads(row['data'])
                if approved:
                    components=json.loads(row['components'])
                    if not components or any(not identifier(component.get('name')) for component in components):
                        raise ValueError('출고 품목과 품목명이 준비된 주문만 강제 출고할 수 있습니다.')
                    source_quantity=quantity(data.get('quantity'))
                    if any(quantity(component.get('quantity')) != source_quantity for component in components):
                        raise ValueError('원본 수량과 출고 품목 수량이 같은 주문만 강제 출고할 수 있습니다.')
                    data['force_shipping_approved']=True
                    data['force_shipping_reason']=reason
                    data['force_shipping_approved_at']=datetime.now().isoformat(timespec='seconds')
                else:
                    for key in ('force_shipping_approved','force_shipping_reason','force_shipping_approved_at'):
                        data.pop(key,None)
                self.db.execute('UPDATE orders SET data=? WHERE id=?',(encode(data),row['id']))
            self.event('강제 출고 승인' if approved else '강제 출고 승인 해제',f'{len(rows)}건'+(f' / {reason}' if approved else ''))
        return len(rows)

    def refresh_matches(self):
        with self.db:
            for order in self.orders():
                if order['state'] not in ('검토 필요','출고 준비'):continue
                components,issue=self.match(order['data'])
                self.db.execute('UPDATE orders SET components=?,state=?,issue=? WHERE id=?',
                    (encode(components),'검토 필요' if issue else '출고 준비',issue,order['id']))

    def orders(self):
        orders = [{**dict(row), 'data':json.loads(row['data']), 'components':json.loads(row['components'])}
                  for row in self.db.execute('SELECT * FROM orders ORDER BY rowid DESC')]
        request_days: dict[str,set[str]] = {}
        shipment_days: dict[str,set[str]] = {}
        for row in self.db.execute('''SELECT r.order_id,q.day FROM request_lines r
            JOIN requests q ON q.id=r.request_id'''):
            request_days.setdefault(row['order_id'],set()).add(row['day'])
        for row in self.db.execute('''SELECT r.order_id,s.day FROM shipments s
            JOIN request_lines r ON r.id=s.line_id'''):
            shipment_days.setdefault(row['order_id'],set()).add(row['day'])
        counts = Counter(self.duplicate_key(order['data']) for order in orders)
        for order in orders:
            order['duplicate_count'] = counts[self.duplicate_key(order['data'])]
            registered=order['data'].get('_registered_day','')
            order['registered_day']=registered
            order['request_days']=sorted(request_days.get(order['id'],set()))
            order['shipment_days']=sorted(shipment_days.get(order['id'],set()))
            order['activity_days']=sorted(set(filter(None,[registered,*order['request_days'],*order['shipment_days']])))
        return sorted(orders, key=lambda order:order['duplicate_count'] <= 1)

    def event_rules(self):
        return [dict(row) for row in self.db.execute('SELECT * FROM event_rules ORDER BY channel,name,source_product,source_option')]

    def set_event_rule(self, order_id, name, target_product, target_option, target_amount, components):
        row = self.db.execute('SELECT * FROM orders WHERE id=?',(order_id,)).fetchone()
        if not row or row['state'] not in ('검토 필요','출고 준비'):
            raise ValueError('출고요청 전 주문에서만 이벤트 규칙을 저장할 수 있습니다.')
        data = json.loads(row['data'])
        source_product, source_option, _ = self._event_source(data)
        name = identifier(name)
        if not name:
            raise ValueError('이벤트 이름을 입력하세요.')
        target_product = identifier(target_product) or source_product
        target_option = identifier(target_option) or source_option
        target_amount = str(number(target_amount)) if identifier(target_amount) else ''
        components, customer = self._prepare_component_definitions(data['channel'], components)
        existing = self.db.execute('''SELECT id FROM event_rules WHERE channel=? AND source_product=? AND source_option=?''',
            (data['channel'],source_product,source_option)).fetchone()
        rule_id = existing['id'] if existing else 'EV-' + uuid.uuid4().hex[:12]
        with self.db:
            self.db.execute('''INSERT OR REPLACE INTO event_rules
                (id,channel,name,source_product,source_option,target_product,target_option,target_amount,components,active)
                VALUES(?,?,?,?,?,?,?,?,?,1)''', (rule_id,data['channel'],name,source_product,source_option,target_product,target_option,target_amount,encode(components)))
            rule = self.db.execute('SELECT * FROM event_rules WHERE id=?',(rule_id,)).fetchone()
            for pending in self.orders():
                current = pending['data']
                current_product, current_option, current_amount = self._event_source(current)
                if pending['state'] not in ('검토 필요','출고 준비') or current['channel'] != data['channel'] or (current_product,current_option) != (source_product,source_option):
                    continue
                current['source_product'], current['source_option'], current['source_amount'] = current_product, current_option, current_amount
                current['product'] = rule['target_product'] or current_product
                current['option'] = rule['target_option'] or current_option
                current['amount'] = str(number(rule['target_amount']) * quantity(current['quantity'])) if rule['target_amount'] else current_amount
                current['event_name'], current['event_rule_id'] = rule['name'], rule_id
                current['event_unit_amount'] = rule['target_amount']
                current['event_components'] = json.loads(rule['components'] or '[]')
                matched, issue = self.match(current)
                self.db.execute('UPDATE orders SET data=?,components=?,state=?,issue=? WHERE id=?',
                    (encode(current),encode(matched),'검토 필요' if issue else '출고 준비',issue,pending['id']))
            self.event('이벤트 규칙 저장', f'{data["channel"]} / {name} / {source_product} / {source_option}')
        if customer:
            self.settings.setdefault('channel_customer_codes',{})[data['channel']] = customer
            self._write_settings()
        return rule_id

    def clear_event_rule(self, order_id):
        row = self.db.execute('SELECT * FROM orders WHERE id=?',(order_id,)).fetchone()
        if not row:
            raise ValueError('주문을 찾지 못했습니다.')
        selected = json.loads(row['data'])
        rule_id = selected.get('event_rule_id')
        if not rule_id:
            source_product, source_option, _ = self._event_source(selected)
            rule = self.db.execute('SELECT id FROM event_rules WHERE channel=? AND source_product=? AND source_option=?',
                (selected['channel'],source_product,source_option)).fetchone()
            rule_id = rule['id'] if rule else None
        if not rule_id:
            raise ValueError('이 주문에 연결된 이벤트 규칙이 없습니다.')
        with self.db:
            self.db.execute('DELETE FROM event_rules WHERE id=?',(rule_id,))
            for pending in self.orders():
                if pending['state'] not in ('검토 필요','출고 준비') or pending['data'].get('event_rule_id') != rule_id:
                    continue
                current = pending['data']
                current['product'] = current.pop('source_product',current['product'])
                current['option'] = current.pop('source_option',current['option'])
                current['amount'] = current.pop('source_amount',current['amount'])
                current.pop('event_name',None); current.pop('event_rule_id',None)
                current.pop('event_components',None); current.pop('event_unit_amount',None)
                matched, issue = self.match(current)
                self.db.execute('UPDATE orders SET data=?,components=?,state=?,issue=? WHERE id=?',
                    (encode(current),encode(matched),'검토 필요' if issue else '출고 준비',issue,pending['id']))
            self.event('이벤트 규칙 삭제',rule_id)

    def delete_orders(self, ids):
        ids = list(dict.fromkeys(ids))
        if not ids:
            raise ValueError('삭제할 주문을 선택하세요.')
        marks = ','.join('?' for _ in ids)
        rows = self.db.execute(f'SELECT id,state FROM orders WHERE id IN ({marks})',ids).fetchall()
        if len(rows) != len(ids):
            raise ValueError('일부 주문을 찾지 못했습니다.')
        with self.db:
            request_rows=self.db.execute(f'SELECT DISTINCT request_id FROM request_lines WHERE order_id IN ({marks})',ids).fetchall()
            request_ids=[row['request_id'] for row in request_rows]
            line_rows=self.db.execute(f'SELECT id FROM request_lines WHERE order_id IN ({marks})',ids).fetchall()
            line_ids=[row['id'] for row in line_rows]
            if line_ids:
                line_marks=','.join('?' for _ in line_ids)
                self.db.execute(f'DELETE FROM shipments WHERE line_id IN ({line_marks})',line_ids)
                self.db.execute(f'DELETE FROM request_lines WHERE id IN ({line_marks})',line_ids)
            self.db.execute(f'DELETE FROM orders WHERE id IN ({marks})',ids)
            for request_id in request_ids:
                self.db.execute("DELETE FROM artifacts WHERE id=? AND kind='출고요청'",(request_id,))
                if not self.db.execute('SELECT 1 FROM request_lines WHERE request_id=? LIMIT 1',(request_id,)).fetchone():
                    self.db.execute('DELETE FROM requests WHERE id=?',(request_id,))
            states=Counter(row['state'] for row in rows)
            self.event('주문 삭제', f'{len(ids)}건 / '+', '.join(f'{state} {count}' for state,count in states.items()))
        return len(ids)

    def set_mapping(self, order_id, components):
        row = self.db.execute('SELECT * FROM orders WHERE id=?', (order_id,)).fetchone()
        if not row or row['state'] not in ('검토 필요', '출고 준비'):
            raise ValueError('요청 전 주문만 매칭을 변경할 수 있습니다.')
        data = json.loads(row['data'])
        components, customer = self._prepare_component_definitions(data['channel'], components)
        key = self.mapping_key(data)
        with self.db:
            self.db.execute('INSERT OR REPLACE INTO mappings VALUES(?,?)', (key, encode(components)))
            for order in self.orders():
                if order['state'] in ('검토 필요', '출고 준비') and self.mapping_key(order['data']) == key:
                    matched, issue = self.match(order['data'])
                    self.db.execute('UPDATE orders SET components=?,state=?,issue=? WHERE id=?', (encode(matched), '검토 필요' if issue else '출고 준비', issue, order['id']))
            self.event('매칭 저장', order_id)
        if customer:
            self.settings.setdefault('channel_customer_codes',{})[data['channel']] = customer
            self._write_settings()

    def edit_delivery(self, order_id, fields):
        with self.db:
            row = self.db.execute('SELECT * FROM orders WHERE id=?', (order_id,)).fetchone()
            if row['state'] not in ('출고 준비', '검토 필요'):
                raise ValueError('출고요청 후 배송정보는 변경할 수 없습니다.')
            data = json.loads(row['data'])
            for key in ('recipient','phone','postcode','address','memo'):
                data[key] = fields[key].strip()
            data = self._normalize_delivery(data)
            components, issue = self.match(data)
            self.db.execute('UPDATE orders SET data=?,components=?,state=?,issue=? WHERE id=?', (encode(data), encode(components), '검토 필요' if issue else '출고 준비', issue, order_id))
            self.event('배송정보 수정', order_id)

    def _artifact(self, artifact_id, kind, on, content, path):
        target = Path(path)
        target.parent.mkdir(parents=True, exist_ok=True)
        fd, temporary = tempfile.mkstemp(suffix='.xlsx', dir=target.parent)
        try:
            with os.fdopen(fd, 'wb') as handle:
                handle.write(content)
            self.db.execute('INSERT INTO artifacts(id,kind,day,content) VALUES(?,?,?,?)', (artifact_id, kind, on, content))
            os.replace(temporary, target)
        finally:
            if os.path.exists(temporary):
                os.unlink(temporary)

    def request(self, ids, on, path, excluded_channels=None):
        on = day(on)
        selected = [o for o in self.orders() if o['id'] in ids]
        if not selected or len(selected) != len(set(ids)):
            raise ValueError('선택한 주문을 다시 확인하세요.')
        orders = [o for o in selected if not is_excluded_channel(o['data'].get('channel'),excluded_channels)]
        if not orders:
            raise ValueError('선택한 주문이 모두 제외 판매처에 포함되어 있습니다.')
        if any(o['state'] != '출고 준비' and not o['data'].get('force_shipping_approved') for o in orders):
            raise ValueError('출고 준비 또는 작업자가 강제 출고 승인한 주문만 선택하세요.')
        for order in orders:
            source_quantity = quantity(order['data']['quantity'])
            if any(quantity(component.get('quantity')) != source_quantity for component in order['components']):
                raise ValueError(
                    f"{order['data']['order_no']}: 최종 출고수량이 원본 엑셀 수량과 다릅니다. "
                    "상품 매칭을 다시 저장한 뒤 출고하세요."
                )
        batch = 'R-' + uuid.uuid4().hex[:12]
        rows = []
        request_format = self.settings.get('request_format', 'wekeep')
        headers = self.settings.get('request_headers', REQUEST_COLUMNS)
        if request_format != 'wekeep' and len(headers) != len(REQUEST_COLUMNS):
            raise ValueError('출고 양식 열 개수가 기본 양식과 다릅니다.')
        with self.db:
            self.db.execute('INSERT INTO requests VALUES(?,?)', (batch, on))
            for order in orders:
                d = order['data']
                for index, c in enumerate(order['components']):
                    line_id = 'L-' + uuid.uuid4().hex
                    self.db.execute('INSERT INTO request_lines VALUES(?,?,?,?,?)', (line_id,batch,order['id'],index,c['quantity']))
                    if request_format == 'wekeep':
                        channel = d['channel'].removeprefix('리큐엠_')
                        rows.append([d['order_no'],channel,c['name'],c['quantity'],d['recipient'],d['phone'],d['postcode'],d['address'],d['memo'],''])
                    else:
                        rows.append([line_id,on,d['channel'],d['account'],d['order_no'],d['line_no'],d['recipient'],d['phone'],d['postcode'],d['address'],d['memo'],c['logistics_code'],c['name'],c['quantity']])
                self.db.execute("UPDATE orders SET state='출고 요청' WHERE id=?", (order['id'],))
            content = wekeep_workbook_bytes(rows) if request_format == 'wekeep' else workbook_bytes(headers, rows, '출고요청')
            self._artifact(batch, '출고요청', on, content, path)
            self.event('출고요청 파일 생성', batch)
        return batch

    def _record_shipment(self, line, shipped_on, shipped_quantity, tracking):
        q = quantity(shipped_quantity)
        order_row = self.db.execute('SELECT data,components FROM orders WHERE id=?', (line['order_id'],)).fetchone()
        if not order_row:
            raise ValueError('출고요청에 연결된 주문을 찾지 못했습니다.')
        order_data = json.loads(order_row['data'])
        if is_smartstore_channel(order_data.get('channel')):
            return 'smartstore'
        old = self.db.execute(
            'SELECT qty FROM shipments WHERE line_id=? AND tracking=? AND day=?',
            (line['id'],tracking,shipped_on),
        ).fetchone()
        if old:
            if old[0] != q:
                raise ValueError('같은 출고 결과의 수량이 변경되었습니다. 이번 입력은 취소됩니다.')
            return 'duplicate'
        done = self.db.execute('SELECT COALESCE(SUM(qty),0) FROM shipments WHERE line_id=?', (line['id'],)).fetchone()[0]
        if done + q > line['qty']:
            raise ValueError('출고수량이 요청수량을 초과합니다. 이번 입력은 취소됩니다.')
        component = json.loads(order_row['components'])[line['component']]
        full = number(component['amount'])
        cumulative = lambda value: (full*Decimal(value)/Decimal(component['quantity'])).quantize(Decimal('1'),rounding=ROUND_HALF_UP)
        suggested_amount = cumulative(done+q)-cumulative(done)
        self.db.execute(
            'INSERT INTO shipments(id,line_id,qty,tracking,day,erp_id,erp_amount) VALUES(?,?,?,?,?,NULL,?)',
            (uuid.uuid4().hex,line['id'],q,tracking,shipped_on,str(suggested_amount)),
        )
        totals = self.db.execute('SELECT SUM(qty) FROM request_lines WHERE order_id=?', (line['order_id'],)).fetchone()[0]
        shipped = self.db.execute(
            'SELECT COALESCE(SUM(s.qty),0) FROM shipments s JOIN request_lines r ON s.line_id=r.id WHERE r.order_id=?',
            (line['order_id'],),
        ).fetchone()[0]
        self.db.execute('UPDATE orders SET state=? WHERE id=?', ('출고 완료' if totals == shipped else '부분 출고',line['order_id']))
        return 'added'

    def auto_import_shipments(self, on, excluded_channels=None):
        """Confirm unprocessed non-SmartStore request lines created on the selected day."""
        shipped_on = day(on)
        rows = self.db.execute('''SELECT r.* FROM request_lines r
            JOIN requests q ON q.id=r.request_id
            WHERE q.day=? ORDER BY r.rowid''', (shipped_on,)).fetchall()
        added = duplicates = smartstore_skipped = excluded_skipped = 0
        with self.db:
            for line in rows:
                if self.db.execute('SELECT 1 FROM shipments WHERE line_id=? LIMIT 1', (line['id'],)).fetchone():
                    duplicates += 1
                    continue
                order_row=self.db.execute('SELECT data FROM orders WHERE id=?',(line['order_id'],)).fetchone()
                order_data=json.loads(order_row['data']) if order_row else {}
                if is_excluded_channel(order_data.get('channel'),excluded_channels):
                    excluded_skipped += 1
                    continue
                result = self._record_shipment(line,shipped_on,line['qty'],'')
                if result == 'added':
                    added += 1
                elif result == 'smartstore':
                    smartstore_skipped += 1
            self.event('당일 출고건 자동 반영',f'신규 {added}, 기존 {duplicates}, 스마트스토어 제외 {smartstore_skipped}, 선택 제외 {excluded_skipped}, 출고일 {shipped_on}')
        return added, duplicates

    def _wekeep_result_line(self, values, headers, shipped_on):
        def value(name):
            if name not in headers:
                return ''
            index=headers.index(name)
            return identifier(values[index]) if index < len(values) else ''
        serial = value('일련번호')
        if serial:
            line = self.db.execute('SELECT * FROM request_lines WHERE id=?', (serial,)).fetchone()
            if not line:
                raise ValueError(f'알 수 없는 일련번호: {serial}')
            return line, value('수량'), value('송장번호')
        order_no, recipient, product = value('주문번호'), value('수령자'), value('상품명')
        candidates = []
        for line in self.db.execute('''SELECT r.* FROM request_lines r
            JOIN requests q ON q.id=r.request_id JOIN orders o ON o.id=r.order_id
            WHERE q.day<=? ORDER BY r.rowid''', (shipped_on,)):
            order = self.db.execute('SELECT data,components FROM orders WHERE id=?', (line['order_id'],)).fetchone()
            data, components = json.loads(order['data']), json.loads(order['components'])
            component = components[line['component']]
            same_order=identifier(data.get('order_no'))==order_no
            same_recipient=identifier(data.get('recipient')).replace(' ','')==recipient.replace(' ','')
            if same_order and same_recipient:
                candidates.append(line)
        product_matches=[]
        for line in candidates:
            order=self.db.execute('SELECT components FROM orders WHERE id=?',(line['order_id'],)).fetchone()
            component=json.loads(order['components'])[line['component']]
            if product and identifier(component.get('name'))==product:product_matches.append(line)
        if product_matches:candidates=product_matches
        requested=quantity(value('수량'))
        available=[]
        for line in candidates:
            shipped=self.db.execute('SELECT COALESCE(SUM(qty),0) FROM shipments WHERE line_id=?',(line['id'],)).fetchone()[0]
            if line['qty']-shipped >= requested:
                available.append(line)
        if not available:
            raise ValueError(f'{order_no} / {recipient}: 출고요청 행을 찾을 수 없거나 이미 전량 반영됐습니다.')
        return available[0], value('수량'), value('송장번호')

    def import_results(self, path, on=None):
        rows = read_rows(path)
        columns = self.settings['result_columns']
        headers = [identifier(x) for x in rows[0]]
        legacy_format = set(columns.values()).issubset(headers)
        wekeep_format = set(WEKEEP_REQUEST_COLUMNS).issubset(headers)
        if not legacy_format and not wekeep_format:
            raise ValueError('물류 결과 헤더가 맞지 않습니다. 결과 양식 또는 설정을 확인하세요.')
        added = duplicates = smartstore_skipped = 0
        affected_orders=[]
        with self.db:
            for row in rows[1:]:
                if not any(x not in (None, '') for x in row):
                    continue
                if wekeep_format:
                    if not on:
                        raise ValueError('위킵 출고파일을 반영할 실제 출고일이 필요합니다.')
                    shipped_on = day(on)
                    line, q, tracking = self._wekeep_result_line(row,headers,shipped_on)
                    line_id = line['id']
                else:
                    d = {key: identifier(row[headers.index(value)]) if headers.index(value) < len(row) else '' for key,value in columns.items()}
                    line_id, tracking, shipped_on = d['요청행ID'], d['송장번호'], day(d['실제출고일'])
                    q = d['출고수량']
                    line = self.db.execute('SELECT * FROM request_lines WHERE id=?', (line_id,)).fetchone()
                if not tracking:
                    raise ValueError('송장번호가 필요합니다.')
                if not line:
                    raise ValueError(f'알 수 없는 요청행ID: {line_id}')
                result = self._record_shipment(line,shipped_on,q,tracking)
                if result == 'smartstore':
                    smartstore_skipped += 1
                    continue
                if result == 'duplicate':
                    duplicates += 1
                    continue
                added += 1
                affected_orders.append(line['order_id'])
            self._sync_source_tracking(affected_orders)
            self.event('물류 결과 반영', f'신규 {added}, 중복 {duplicates}, 스마트스토어 제외 {smartstore_skipped}')
        return added, duplicates

    def _smartstore_erp_profile(self, purpose='smartstore_erp'):
        return next((profile for profile in self.settings.get('profiles', [])
                     if profile.get('purpose') == purpose or (
                         purpose == 'smartstore_erp' and profile.get('channel') == SMARTSTORE_ERP_CHANNEL
                     )), None)

    def smartstore_erp_enabled(self):
        return any(
            profile and profile.get('enabled') and profile.get('columns')
            for profile in (
                self._smartstore_erp_profile('smartstore_erp'),
                self._smartstore_erp_profile('smartstore_purchase_erp'),
            )
        )

    @staticmethod
    def _reprice_components(components, total, source_quantity):
        rows = [dict(component) for component in components]
        if not rows:
            return []
        total = number(total)
        weights = [number(component.get('amount','0')) for component in rows]
        weight_total = sum(weights, Decimal('0'))
        allocated = []
        remaining = total
        for index, component in enumerate(rows):
            if index == 0:
                amount = Decimal('0')
            elif weight_total:
                amount = (total * weights[index] / weight_total).quantize(Decimal('1'), rounding=ROUND_HALF_UP)
                remaining -= amount
            else:
                amount = Decimal('0')
            allocated.append({**component, 'quantity':source_quantity, 'source_quantity':source_quantity, 'amount':str(amount)})
        allocated[0]['amount'] = str(remaining)
        return allocated

    def _smartstore_erp_components(self, data):
        candidates = [order for order in self.orders() if is_smartstore_channel(order['data'].get('channel'))]
        original = next((order for order in candidates if data.get('line_no') and order['data'].get('line_no') == data['line_no']), None)
        if original is None:
            original = next((order for order in candidates if order['data'].get('order_no') == data.get('order_no')
                             and order['data'].get('product') == data.get('product')
                             and order['data'].get('option') == data.get('option')), None)
        if original and original['components']:
            return self._reprice_components(original['components'], data['amount'], quantity(data['quantity'])), ''
        lookup = {**data, 'channel':SMARTSTORE_CHANNEL, 'account':data.get('account','기본')}
        custom = self.db.execute('SELECT components FROM mappings WHERE key=?', (self.mapping_key(lookup),)).fetchone()
        if custom:
            components, issue = self._build_components(lookup, json.loads(custom[0]))
            return components, issue
        return [], '스마트스토어 주문 또는 저장된 상품·세트 매칭을 찾지 못했습니다.'

    @staticmethod
    def _smartstore_rows_match(left, right):
        if left.get('line_no') and right.get('line_no'):
            return left.get('line_no') == right.get('line_no')
        return (
            left.get('order_no') == right.get('order_no')
            and left.get('product') == right.get('product')
            and left.get('option','') == right.get('option','')
        )

    def _find_smartstore_erp_row(self, data, *, pending_only=False):
        query='SELECT * FROM smartstore_erp_rows'
        if pending_only:
            query += ' WHERE erp_id IS NULL'
        for row in self.db.execute(query+' ORDER BY rowid'):
            if self._smartstore_rows_match(json.loads(row['data']),data):
                return row
        return None

    def import_smartstore_erp(self, path, on):
        profile = self._smartstore_erp_profile('smartstore_erp')
        if not profile or not profile.get('enabled'):
            raise ValueError('매칭 설정에서 스마트스토어 ERP매칭 파일과 열을 먼저 저장하세요.')
        shipped_on = day(on)
        incoming = parse_orders(path, [profile], channel_override=SMARTSTORE_ERP_CHANNEL)
        added = duplicates = 0
        with self.db:
            for data in incoming:
                data['channel'] = SMARTSTORE_CHANNEL
                data['account'] = data.get('account') or '기본'
                data['quantity'] = str(quantity(data.get('quantity') or '1'))
                data['amount'] = str(money(data.get('amount') or '0'))
                data['shipping'] = str(money(data.get('shipping') or '0'))
                token = encode([
                    shipped_on, data.get('order_no'), data.get('line_no'), data.get('product'),
                    data.get('option'), data.get('quantity'), data.get('amount'),
                ])
                identity = hashlib.sha256(token.encode('utf-8')).hexdigest()
                if self.db.execute('SELECT 1 FROM smartstore_erp_rows WHERE identity=?',(identity,)).fetchone():
                    duplicates += 1
                    continue
                if any(
                    json.loads(row['data']).get('_shipping_import_identity') == identity
                    for row in self.db.execute('SELECT data FROM smartstore_erp_rows')
                ):
                    duplicates += 1
                    continue
                purchase_target=self._find_smartstore_erp_row(data,pending_only=True)
                previous=json.loads(purchase_target['data']) if purchase_target else {}
                if purchase_target and previous.get('_purchase_only'):
                    if previous.get('_shipping_import_identity') == identity:
                        duplicates += 1
                        continue
                    confirmed=bool(previous.get('_purchase_confirmed'))
                    merged={**data,**({
                        key:previous.get(key) for key in ('amount','quantity','paid_at') if previous.get(key) not in (None,'')
                    } if confirmed else {})}
                    merged['_purchase_confirmed']=confirmed
                    merged['_purchase_only']=False
                    merged['_shipping_import_identity']=identity
                    old_components=json.loads(purchase_target['components'])
                    if old_components:
                        components=self._reprice_components(old_components,merged['amount'],quantity(merged['quantity']))
                        issue=''
                    else:
                        components,issue=self._smartstore_erp_components(merged)
                    self.db.execute('UPDATE smartstore_erp_rows SET data=?,components=?,issue=? WHERE id=?',(
                        encode(merged),encode(components),issue,purchase_target['id'],
                    ))
                    added += 1
                    continue
                data['_shipping_import_identity']=identity
                data['_purchase_confirmed']=False
                data['_purchase_only']=False
                components, issue = self._smartstore_erp_components(data)
                self.db.execute('INSERT INTO smartstore_erp_rows VALUES(?,?,?,?,?,?,?)',
                    (uuid.uuid4().hex,identity,shipped_on,encode(data),encode(components),None,issue))
                added += 1
            self.event('스마트스토어 ERP 원본 반영',f'신규 {added}, 중복 {duplicates}, 출고일 {shipped_on}')
        return added, duplicates

    def import_smartstore_purchase(self, path, on):
        profile=self._smartstore_erp_profile('smartstore_purchase_erp')
        if not profile or not profile.get('enabled'):
            raise ValueError('매칭 설정에서 스마트스토어 ERP매칭(구매확정파일)의 파일과 열을 먼저 저장하세요.')
        confirmed_on=day(on)
        incoming=parse_orders(path,[profile],channel_override=SMARTSTORE_PURCHASE_CHANNEL)
        added=duplicates=0
        with self.db:
            for data in incoming:
                data['channel']=SMARTSTORE_CHANNEL
                data['account']=data.get('account') or '기본'
                data['quantity']=str(quantity(data.get('quantity') or '1'))
                data['amount']=str(money(data.get('amount') or '0'))
                data['shipping']=str(money(data.get('shipping') or '0'))
                token=encode([
                    data.get('order_no'),data.get('line_no'),data.get('product'),data.get('option'),
                    data.get('quantity'),data.get('amount'),data.get('paid_at'),
                ])
                identity=hashlib.sha256(token.encode('utf-8')).hexdigest()
                if self.db.execute('SELECT 1 FROM smartstore_purchase_rows WHERE identity=?',(identity,)).fetchone():
                    duplicates += 1
                    continue
                target=self._find_smartstore_erp_row(data,pending_only=True)
                if target:
                    previous=json.loads(target['data'])
                    merged={**previous,**data,'_purchase_confirmed':True,'_purchase_only':False}
                    old_components=json.loads(target['components'])
                    if old_components:
                        components=self._reprice_components(old_components,merged['amount'],quantity(merged['quantity']))
                        issue=''
                    else:
                        components,issue=self._smartstore_erp_components(merged)
                    row_id=target['id']
                    self.db.execute('UPDATE smartstore_erp_rows SET data=?,components=?,issue=? WHERE id=?',(
                        encode(merged),encode(components),issue,row_id,
                    ))
                else:
                    data.update(_purchase_confirmed=True,_purchase_only=True)
                    components,issue=self._smartstore_erp_components(data)
                    row_id=uuid.uuid4().hex
                    row_identity=hashlib.sha256(('purchase:'+identity).encode('utf-8')).hexdigest()
                    self.db.execute('INSERT INTO smartstore_erp_rows VALUES(?,?,?,?,?,?,?)',(
                        row_id,row_identity,confirmed_on,encode(data),encode(components),None,issue,
                    ))
                self.db.execute('INSERT INTO smartstore_purchase_rows VALUES(?,?,?,?,?)',(
                    uuid.uuid4().hex,identity,confirmed_on,encode(data),row_id,
                ))
                added += 1
            self.event('스마트스토어 구매확정 원본 반영',f'신규 {added}, 중복 {duplicates}, 기준일 {confirmed_on}')
        return added,duplicates

    def _erp_source_components(self, data):
        custom = self.db.execute('SELECT components FROM mappings WHERE key=?', (self.mapping_key(data),)).fetchone()
        if custom:
            return self._build_components(data, json.loads(custom[0]))
        catalog_channel = self.catalog_channel(data['channel'])
        if catalog_channel not in self.catalog.channels:
            return [], f'{data["channel"]}의 ERP 거래처코드가 없습니다.'
        order = SmartStoreOrder(
            0, data.get('order_no',''), data.get('order_no',''), datetime.now(), '결제완료',
            data.get('product',''), data.get('option',''), Decimal(data['quantity']),
            number(data['amount']), include_shipping=False,
        )
        converted = convert_orders(
            [order], self.conversion_catalog(catalog_channel), channel_name=catalog_channel,
            default_warehouse=self.settings['warehouse'],
        )
        if converted.issues or not converted.is_reconciled:
            reason = '; '.join(issue.reason for issue in converted.issues) or 'ERP 상품 변환 규칙이 없습니다.'
            return [], reason
        grouped = {}
        for line in converted.lines:
            if not line.item_code or line.item_code not in self.catalog.items:
                return [], '등록된 ERP 품목코드가 필요합니다.'
            key = (line.item_code, line.warehouse)
            if key not in grouped:
                grouped[key] = dict(
                    code=line.item_code, logistics_code=line.item_code, name=line.item_name,
                    quantity=0, amount='0', warehouse=line.warehouse, customer=line.customer_code,
                    match_method='판매전표 DB 변환 규칙',
                )
            component = grouped[key]
            component['quantity'] += quantity(line.quantity)
            component['amount'] = str(number(component['amount']) + line.total)
        return list(grouped.values()), ''

    def import_esm_erp(self, paths, on):
        imported_on = day(on)
        if isinstance(paths, (str, Path)):
            paths = [paths]
        added = duplicates = 0
        with self.db:
            for path in paths:
                for data in parse_esm_rows(path):
                    data['account'] = data.get('account') or '기본'
                    row_day = day(data.get('order_day') or imported_on)
                    token = encode([
                        data.get('channel'), data.get('account'), data.get('order_no'), data.get('product'),
                        data.get('source_product_no'), data.get('option'), data.get('quantity'), data.get('amount'),
                    ])
                    identity = hashlib.sha256(token.encode('utf-8')).hexdigest()
                    if self.db.execute('SELECT 1 FROM esm_erp_rows WHERE identity=?',(identity,)).fetchone():
                        duplicates += 1
                        continue
                    components, issue = self._erp_source_components(data)
                    self.db.execute('INSERT INTO esm_erp_rows VALUES(?,?,?,?,?,?,?)',(
                        uuid.uuid4().hex, identity, row_day, encode(data), encode(components), None, issue,
                    ))
                    added += 1
            self.event('옥션/지마켓 ERP 원본 반영', f'신규 {added}, 중복 {duplicates}, 기준일 {imported_on}')
        return added, duplicates

    def confirmed_erp_entries(self, on=None):
        selected_day = day(on) if on else None
        entries = []
        query = '''SELECT s.*, r.order_id, r.component FROM shipments s
            JOIN request_lines r ON s.line_id=r.id'''
        parameters = []
        if selected_day:
            query += ' WHERE s.day=?'; parameters.append(selected_day)
        orders = {order['id']:order for order in self.orders()}
        for row in self.db.execute(query+' ORDER BY s.day,s.rowid',parameters):
            order = orders[row['order_id']]
            if self.smartstore_erp_enabled() and is_smartstore_channel(order['data'].get('channel')):
                continue
            component = order['components'][row['component']]
            entries.append({'source':'일반 실제출고','id':row['id'],'day':row['day'],
                'channel':order['data']['channel'],'order_no':order['data']['order_no'],
                'recipient':order['data']['recipient'],'product':component.get('name') or component.get('code',''),
                'code':component.get('code',''),'quantity':row['qty'],'amount':row['erp_amount'] or '',
                'tracking':row['tracking'],'issue':'','erp_id':row['erp_id']})
        query = 'SELECT * FROM smartstore_erp_rows'
        parameters = []
        if selected_day:
            query += ' WHERE day=?';parameters.append(selected_day)
        for row in self.db.execute(query+' ORDER BY day,rowid',parameters):
            data=json.loads(row['data']);components=json.loads(row['components'])
            if not components:
                entries.append({'source':'스마트스토어 ERP','id':f"SSERP:{row['id']}:0",'day':row['day'],
                    'channel':'스마트스토어','order_no':data.get('order_no',''),'recipient':data.get('recipient',''),
                    'product':data.get('product',''),'code':'','quantity':data.get('quantity',''),
                    'amount':'','tracking':'','issue':row['issue'],'erp_id':row['erp_id']})
            for index,component in enumerate(components):
                entries.append({'source':'스마트스토어 ERP','id':f"SSERP:{row['id']}:{index}",'day':row['day'],
                    'channel':'스마트스토어','order_no':data.get('order_no',''),'recipient':data.get('recipient',''),
                    'product':component.get('name') or component.get('code',''),'code':component.get('code',''),
                    'quantity':component.get('quantity',data.get('quantity','')),'amount':component.get('amount',''),
                    'tracking':'','issue':row['issue'],'erp_id':row['erp_id']})
        return entries

    def pending_erp_shipments(self, on=None):
        """Return actual shipments waiting for ERP, including their editable ERP amount."""
        parameters = []
        condition = 's.erp_id IS NULL'
        if on:
            condition += ' AND s.day=?'
            parameters.append(day(on))
        rows = self.db.execute(f'''SELECT s.*, r.order_id, r.component
            FROM shipments s JOIN request_lines r ON s.line_id=r.id
            WHERE {condition} ORDER BY s.day,s.rowid''', parameters).fetchall()
        orders = {order['id']:order for order in self.orders()}
        result = []
        for row in rows:
            order = orders[row['order_id']]
            data = order['data']
            if self.smartstore_erp_enabled() and is_smartstore_channel(data.get('channel')):
                continue
            component = order['components'][row['component']]
            result.append({
                'id':row['id'], 'day':row['day'], 'order_no':data['order_no'],
                'recipient':data['recipient'], 'product':component.get('name') or component.get('code',''),
                'code':component.get('code',''), 'quantity':row['qty'],
                'amount':row['erp_amount'] or '', 'tracking':row['tracking'], 'source':'일반 실제출고',
            })
        return result

    def pending_erp_entries(self, on=None):
        return [entry for entry in self.confirmed_erp_entries(on) if not entry.get('erp_id')]

    def set_erp_amount(self, shipment_id, amount):
        value = number(amount)
        if value != value.to_integral_value():
            raise ValueError('ERP 반영 금액은 원 단위 정수로 입력하세요.')
        if str(shipment_id).startswith('SSERP:'):
            _,row_id,index_text = shipment_id.split(':',2)
            with self.db:
                row=self.db.execute('SELECT components,erp_id FROM smartstore_erp_rows WHERE id=?',(row_id,)).fetchone()
                if not row:raise ValueError('스마트스토어 ERP 내역을 찾지 못했습니다.')
                if row['erp_id']:raise ValueError('이미 ERP 파일에 반영된 금액은 변경할 수 없습니다.')
                components=json.loads(row['components']);index=int(index_text)
                if not (0 <= index < len(components)):raise ValueError('세트 구성품을 먼저 설정하세요.')
                components[index]['amount']=str(value)
                self.db.execute('UPDATE smartstore_erp_rows SET components=? WHERE id=?',(encode(components),row_id))
                self.event('ERP 금액 매칭',f'{shipment_id}: {value}')
            return
        with self.db:
            row = self.db.execute('SELECT erp_id FROM shipments WHERE id=?', (shipment_id,)).fetchone()
            if not row:
                raise ValueError('실제 출고 내역을 찾지 못했습니다.')
            if row['erp_id']:
                raise ValueError('이미 ERP 파일에 반영된 출고 건은 금액을 변경할 수 없습니다.')
            self.db.execute('UPDATE shipments SET erp_amount=? WHERE id=?', (str(value),shipment_id))
            self.event('ERP 금액 매칭', f'{shipment_id}: {value}')

    def smartstore_erp_row(self, entry_id):
        if not str(entry_id).startswith('SSERP:'):
            return None
        row_id=entry_id.split(':',2)[1]
        row=self.db.execute('SELECT * FROM smartstore_erp_rows WHERE id=?',(row_id,)).fetchone()
        return {**dict(row),'data':json.loads(row['data']),'components':json.loads(row['components'])} if row else None

    def set_smartstore_erp_components(self, entry_id, components):
        row=self.smartstore_erp_row(entry_id)
        if not row:raise ValueError('스마트스토어 ERP 행에서만 세트 구성을 수정할 수 있습니다.')
        if row['erp_id']:raise ValueError('이미 ERP 파일로 만든 행은 수정할 수 없습니다.')
        definitions,customer=self._prepare_component_definitions(SMARTSTORE_CHANNEL,components)
        built,issue=self._build_components(row['data'],definitions)
        if issue:raise ValueError(issue)
        with self.db:
            self.db.execute('UPDATE smartstore_erp_rows SET components=?,issue=? WHERE id=?',(encode(built),'',row['id']))
            self.db.execute('INSERT OR REPLACE INTO mappings VALUES(?,?)',(self.mapping_key(row['data']),encode(definitions)))
            self.event('스마트스토어 ERP 세트 분리',row['id'])
        if customer:
            self.settings.setdefault('channel_customer_codes',{})[SMARTSTORE_CHANNEL]=customer
            self._write_settings()

    def esm_erp_entries(self, on=None, pending_only=False):
        conditions, parameters = [], []
        if on:
            conditions.append('day=?'); parameters.append(day(on))
        if pending_only:
            conditions.append('erp_id IS NULL')
        query = 'SELECT * FROM esm_erp_rows'
        if conditions:
            query += ' WHERE ' + ' AND '.join(conditions)
        entries = []
        for row in self.db.execute(query + ' ORDER BY day,rowid', parameters):
            data, components = json.loads(row['data']), json.loads(row['components'])
            entries.append({
                'id':f"ESM:{row['id']}", 'day':row['day'], 'channel':data.get('channel',''),
                'order_no':data.get('order_no',''), 'product':data.get('product',''),
                'option':data.get('option',''), 'quantity':data.get('quantity',''),
                'unit_amount':data.get('unit_amount',''), 'amount':data.get('amount',''),
                'component_summary':' / '.join(component.get('code','') for component in components),
                'issue':row['issue'], 'erp_id':row['erp_id'],
            })
        return entries

    def esm_erp_row(self, entry_id):
        if not str(entry_id).startswith('ESM:'):
            return None
        row = self.db.execute('SELECT * FROM esm_erp_rows WHERE id=?',(str(entry_id).split(':',1)[1],)).fetchone()
        return {**dict(row),'data':json.loads(row['data']),'components':json.loads(row['components'])} if row else None

    def set_esm_erp_components(self, entry_id, components):
        row = self.esm_erp_row(entry_id)
        if not row:
            raise ValueError('옥션/지마켓 ERP 행을 찾지 못했습니다.')
        if row['erp_id']:
            raise ValueError('이미 ERP 파일로 만든 행은 수정할 수 없습니다.')
        channel = row['data']['channel']
        definitions, customer = self._prepare_component_definitions(channel, components)
        built, issue = self._build_components(row['data'], definitions)
        if issue:
            raise ValueError(issue)
        with self.db:
            self.db.execute('INSERT OR REPLACE INTO mappings VALUES(?,?)',(self.mapping_key(row['data']),encode(definitions)))
            for candidate in self.db.execute('SELECT id,data FROM esm_erp_rows WHERE erp_id IS NULL').fetchall():
                candidate_data=json.loads(candidate['data'])
                if self.mapping_key(candidate_data) != self.mapping_key(row['data']):
                    continue
                candidate_components,candidate_issue=self._build_components(candidate_data,definitions)
                self.db.execute('UPDATE esm_erp_rows SET components=?,issue=? WHERE id=?',(
                    encode(candidate_components),candidate_issue,candidate['id'],
                ))
            self.event('옥션/지마켓 ERP 상품 매칭', row['id'])
        if customer:
            self.settings.setdefault('channel_customer_codes',{})[channel] = customer
            self._write_settings()

    def export_erp(self, on, through, path, from_day=None, include_esm=True, only_esm=False,
                   excluded_channels=None):
        on, through = day(on), day(through)
        from_day = day(from_day) if from_day else '0001-01-01'
        if from_day > through:
            raise ValueError('출고 시작일은 종료일보다 늦을 수 없습니다.')
        batch = 'E-' + uuid.uuid4().hex[:12]
        pending = self.db.execute('''SELECT s.*, r.order_id, r.component FROM shipments s
            JOIN request_lines r ON s.line_id=r.id
            WHERE s.erp_id IS NULL AND s.day>=? AND s.day<=? ORDER BY s.day,s.rowid''',
            (from_day,through)).fetchall()
        orders = {o['id']:o for o in self.orders()}
        pending = [shipment for shipment in pending if not is_excluded_channel(
            orders[shipment['order_id']]['data'].get('channel'),excluded_channels
        )]
        if self.smartstore_erp_enabled():
            pending = [shipment for shipment in pending if not is_smartstore_channel(orders[shipment['order_id']]['data'].get('channel'))]
        smart_rows=self.db.execute('''SELECT * FROM smartstore_erp_rows
            WHERE erp_id IS NULL AND day>=? AND day<=? ORDER BY day,rowid''',(from_day,through)).fetchall()
        esm_rows=self.db.execute('''SELECT * FROM esm_erp_rows
            WHERE erp_id IS NULL AND day>=? AND day<=? ORDER BY day,rowid''',(from_day,through)).fetchall()
        if is_excluded_channel('스마트스토어',excluded_channels):
            smart_rows=[]
        esm_rows=[row for row in esm_rows if not is_excluded_channel(
            json.loads(row['data']).get('channel'),excluded_channels
        )]
        if only_esm:
            pending,smart_rows=[],[]
        elif not include_esm:
            esm_rows=[]
        if not pending and not smart_rows and not esm_rows:
            raise ValueError('선택한 출고일 범위에 ERP 파일로 만들 미반영 주문이 없습니다.')
        unmatched = [shipment for shipment in pending if shipment['erp_amount'] in (None,'')]
        if unmatched:
            raise ValueError(f'ERP 금액 매칭이 필요한 실제 출고가 {len(unmatched):,}건 있습니다.')
        smart_unmatched=[]
        for row in smart_rows:
            components=json.loads(row['components'])
            if row['issue'] or not components or any(component.get('amount') in (None,'') for component in components):
                smart_unmatched.append(row)
        if smart_unmatched:
            raise ValueError(f'금액 매칭 또는 세트 분리가 필요한 스마트스토어 ERP 주문이 {len(smart_unmatched):,}건 있습니다.')
        esm_unmatched=[]
        for row in esm_rows:
            components=json.loads(row['components'])
            if row['issue'] or not components or any(component.get('amount') in (None,'') for component in components):
                esm_unmatched.append(row)
        if esm_unmatched:
            raise ValueError(f'상품 매칭이 필요한 옥션/지마켓 ERP 주문이 {len(esm_unmatched):,}건 있습니다.')
        rows = []
        def append(c, q, amount):
            # Split integer won totals into two unit prices to retain exact totals.
            total = int(amount.quantize(Decimal('1'), rounding=ROUND_HALF_UP))
            price, remainder = divmod(total, q)
            for count, unit in ((q-remainder,price),(remainder,price+1)):
                if count:
                    supply = int((Decimal(unit*count)/Decimal('1.1')).quantize(Decimal('1'), rounding=ROUND_HALF_UP))
                    rows.append([None,None,c['customer'],None,self.settings['manager_code'],int(warehouse_code(c['warehouse'])),None,None,None,None,None,None,c['code'],None,None,count,unit,None,supply,unit*count-supply,None,None])
        with self.db:
            for shipment in pending:
                order = orders[shipment['order_id']]
                c = order['components'][shipment['component']]
                append(c, shipment['qty'], number(shipment['erp_amount']))
                self.db.execute('UPDATE shipments SET erp_id=? WHERE id=?', (batch,shipment['id']))
            for smart_row in smart_rows:
                for component in json.loads(smart_row['components']):
                    append(component,quantity(component['quantity']),number(component['amount']))
                self.db.execute('UPDATE smartstore_erp_rows SET erp_id=? WHERE id=?',(batch,smart_row['id']))
            for esm_row in esm_rows:
                for component in json.loads(esm_row['components']):
                    append(component,quantity(component['quantity']),number(component['amount']))
                self.db.execute('UPDATE esm_erp_rows SET erp_id=? WHERE id=?',(batch,esm_row['id']))
            included_order_ids={shipment['order_id'] for shipment in pending}
            included_bundles={self.bundle_key(orders[order_id]['data']) for order_id in included_order_ids}
            groups = {}
            for order in orders.values():
                if is_excluded_channel(order['data'].get('channel'),excluded_channels):
                    continue
                bundle=self.bundle_key(order['data'])
                if bundle in included_bundles:
                    groups.setdefault(bundle,[]).append(order)
            for bundle, group in groups.items():
                if self.smartstore_erp_enabled() and any(is_smartstore_channel(order['data'].get('channel')) for order in group):
                    continue
                if any(o['state'] != '출고 완료' for o in group):
                    continue
                if any(self.db.execute('SELECT COUNT(*) FROM shipments s JOIN request_lines r ON s.line_id=r.id WHERE r.order_id=? AND s.erp_id IS NULL', (o['id'],)).fetchone()[0] for o in group):
                    continue
                if self.db.execute('SELECT 1 FROM fees WHERE bundle=?',(bundle,)).fetchone():
                    continue
                amounts = {number(o['data']['shipping']) for o in group}
                if len(amounts) != 1:
                    raise ValueError('동일 배송비 묶음의 금액이 서로 다릅니다. ERP 출력을 중단합니다.')
                fee = amounts.pop()
                if fee:
                    base = group[0]['components'][0]
                    append({**base,'code':'택배운송비','name':'배송비'},1,fee)
                self.db.execute('INSERT INTO fees VALUES(?,?)',(bundle,batch))
            rows.sort(key=lambda row: (0 if row[5] == 100 else 1))
            headers = ['일자','순번','거래처코드','거래처명','담당자','출하창고','거래유형','통화','환율','계좌번호','미수금','특이사항','품목코드','품목명','규격','수량','단가','외화금액','공급가액','부가세','비고','생산전표생성']
            kind='ESM ERP' if only_esm else 'ERP'
            self._artifact(batch,kind,on,workbook_bytes(headers,rows,'이카운트 웹입력'),path)
            self.event(f'{kind} 파일 생성',batch)
        return batch

    def artifacts(self):
        return [dict(r) for r in self.db.execute('SELECT id,kind,day,registered FROM artifacts ORDER BY rowid DESC')]

    def artifact_history_rows(self, on=None):
        """Return database-style order details for generated files on one voucher day."""
        params=()
        where=''
        if on:
            where='WHERE day=?';params=(day(on),)
        artifacts=self.db.execute(
            f'SELECT id,kind,day,registered FROM artifacts {where} ORDER BY rowid DESC',params
        ).fetchall()
        orders={row['id']:row for row in self.orders()}
        result=[]

        def add(artifact,source,source_day,data,qty,amount,row_id,tracking='',warehouse=''):
            note=self.db.execute('SELECT note FROM artifact_notes WHERE artifact_id=? AND row_key=?',(
                artifact['id'],str(row_id),
            )).fetchone()
            result.append({
                'row_id':f"{artifact['id']}:{source}:{row_id}",
                'artifact_id':artifact['id'],'artifact_day':artifact['day'],
                'kind':artifact['kind'],'source':source,'source_day':source_day,
                'channel':data.get('channel',''),'order_no':data.get('order_no',''),
                'recipient':data.get('recipient',''),'phone':data.get('phone',''),
                'product':data.get('product',''),'option':data.get('option',''),
                'quantity':str(qty or ''),'amount':str(amount or '0'),
                'tracking':tracking or data.get('tracking',''),'warehouse':warehouse,
                'note':note['note'] if note else '','note_key':str(row_id),
                'source_file':data.get('source_file',''),
                'registered':bool(artifact['registered']),
            })

        for artifact in artifacts:
            if artifact['kind']=='출고요청':
                lines=self.db.execute(
                    'SELECT * FROM request_lines WHERE request_id=? ORDER BY rowid',(artifact['id'],)
                ).fetchall()
                for line in lines:
                    order=orders.get(line['order_id'])
                    if not order:continue
                    components=order['components'];index=line['component']
                    component=components[index] if 0<=index<len(components) else {}
                    tracking=self._tracking_text(value[0] for value in self.db.execute(
                        'SELECT tracking FROM shipments WHERE line_id=? ORDER BY rowid',(line['id'],)
                    ))
                    add(artifact,'출고요청',artifact['day'],order['data'],line['qty'],component.get('amount','0'),
                        line['id'],tracking,warehouse_code(component.get('warehouse','')))
                continue
            shipments=self.db.execute('''SELECT s.*,r.order_id,r.component FROM shipments s
                JOIN request_lines r ON s.line_id=r.id WHERE s.erp_id=? ORDER BY s.day,s.rowid''',
                (artifact['id'],)).fetchall()
            for shipment in shipments:
                order=orders.get(shipment['order_id'])
                if order:
                    component=order['components'][shipment['component']]
                    add(artifact,'일반 ERP',shipment['day'],order['data'],shipment['qty'],shipment['erp_amount'],
                        shipment['id'],shipment['tracking'],warehouse_code(component.get('warehouse','')))
            for table,source in (('smartstore_erp_rows','스마트스토어 ERP'),('esm_erp_rows','ESM ERP')):
                for row in self.db.execute(f'SELECT * FROM {table} WHERE erp_id=? ORDER BY day,rowid',(artifact['id'],)).fetchall():
                    data=json.loads(row['data'])
                    components=json.loads(row['components'])
                    warehouses=self._tracking_text(warehouse_code(component.get('warehouse','')) for component in components)
                    add(artifact,source,row['day'],data,data.get('quantity',''),data.get('amount','0'),row['id'],
                        data.get('tracking',''),warehouses)
        return result

    def set_artifact_note(self, artifact_id, row_key, note):
        with self.db:
            if not self.db.execute('SELECT 1 FROM artifacts WHERE id=?',(artifact_id,)).fetchone():
                raise ValueError('출력 이력을 찾지 못했습니다.')
            self.db.execute('INSERT OR REPLACE INTO artifact_notes(artifact_id,row_key,note) VALUES(?,?,?)',(
                artifact_id,str(row_key),identifier(note),
            ))
            self.event('출력 이력 비고 수정',f'{artifact_id}:{row_key}')

    def reexport(self, artifact_id, path):
        row = self.db.execute('SELECT content FROM artifacts WHERE id=?',(artifact_id,)).fetchone()
        if not row:
            raise ValueError('출력 이력을 찾지 못했습니다.')
        Path(path).write_bytes(row[0])

    def mark_registered(self, artifact_id):
        with self.db:
            if not self.db.execute("SELECT 1 FROM artifacts WHERE id=? AND kind IN ('ERP','ESM ERP')",(artifact_id,)).fetchone():
                raise ValueError('ERP 파일 이력을 선택하세요.')
            self.db.execute('UPDATE artifacts SET registered=1 WHERE id=?',(artifact_id,))
            self.event('ERP 등록 확인',artifact_id)

    def backup(self, path):
        if Path(path).resolve() == (self.folder/'operations.sqlite3').resolve():
            raise ValueError('현재 사용 중인 DB에는 백업할 수 없습니다.')
        destination = sqlite3.connect(path)
        try:
            self.db.backup(destination)
        finally:
            destination.close()
